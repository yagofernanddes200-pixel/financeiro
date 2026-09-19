"""
relatorios.py — Geração dos relatórios financeiros em PDF e Word (mensal e anual).

Extraído de app.py para reduzir o tamanho do arquivo principal. Assim como core.py,
essas funções não dependem do Streamlit: recebem o dicionário `dados` e alguns
parâmetros (período, ano, nome do usuário) e devolvem os bytes do arquivo pronto
(PDF ou DOCX), prontos para serem passados a `st.download_button`.
"""

from datetime import datetime

from core import pertence_periodo_cartao, emprestimo_pendente_no_periodo


# --- FUNÇÕES PARA GERAÇÃO DE RELATÓRIOS (PDF & WORD) ---
def gerar_relatorio_mensal_pdf(dados, periodo, mes_nome, ano, usuario_atual):
    from fpdf import FPDF
    from datetime import datetime
    
    # Filtrar dados do mês selecionado
    receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo)]
    gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if pertence_periodo_cartao(g, periodo)]
    gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if pertence_periodo_cartao(g, periodo)]
    
    total_receitas = sum(r["valor"] for r in receitas_mes)
    total_fixos = sum(g["valor"] for g in gastos_fixos_mes)
    total_avulsos = sum(g["valor"] for g in gastos_avulsos_mes)
    total_gastos = total_fixos + total_avulsos
    saldo_livre = total_receitas - total_gastos
    
    # Economias do mês
    economias_mes = 0.0
    for t in dados.get("transferencias", []):
        if t["data"].startswith(periodo):
            dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
            orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
            if dest_tipo == "Guardado" and orig_tipo == "Normal":
                economias_mes += t["valor"]
                
    # Terceiros pendentes
    unpaid_cartao = [t for t in dados.get("gastos_terceiros_cartao", []) if not t["pago"]]
    unpaid_emprestimo = [e for e in dados.get("gastos_terceiros_emprestimo", []) if not e["pago"]]
    
    total_cartao_terceiros = sum(t["valor"] for t in unpaid_cartao if pertence_periodo_cartao(t, periodo))
    total_emprestimos_pendentes = sum(e["valor"] for e in dados.get("gastos_terceiros_emprestimo", []) if emprestimo_pendente_no_periodo(e, periodo))
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_margins(15, 15, 15)
    pdf.set_text_color(32, 32, 36)
    
    # Cabeçalho
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "FINANCAS DIDATICAS - RELATORIO MENSAL", ln=True, align="C")
    pdf.set_font("Helvetica", "", 12)
    pdf.cell(0, 8, f"Mes de Referencia: {mes_nome} de {ano}", ln=True, align="C")
    pdf.cell(0, 6, f"Usuario: {usuario_atual} | Gerado em: {datetime.now().strftime('%d/%m/%Y as %H:%M')}", ln=True, align="C")
    pdf.ln(8)
    
    # Linha divisória Roxo
    pdf.set_draw_color(130, 87, 229)
    pdf.set_line_width(1)
    pdf.line(15, pdf.get_y(), 195, pdf.get_y())
    pdf.ln(6)
    
    # 1. Resumo Explicativo
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "1. RESUMO EXPLICATIVO DO PERIODO", ln=True)
    pdf.set_font("Helvetica", "", 10.5)
    
    texto = (
        f"Durante o mes de {mes_nome} de {ano}, o perfil '{usuario_atual}' apresentou as seguintes movimentacoes:\n\n"
        f"- Receitas Totais: Entrou um total de R$ {total_receitas:,.2f} em receitas cadastradas.\n"
        f"- Custos Totais: O total de despesas reais pagas ou a pagar no mes somou R$ {total_gastos:,.2f}, dividindo-se em R$ {total_fixos:,.2f} de gastos fixos/parcelamentos e R$ {total_avulsos:,.2f} de gastos avulsos (dia a dia).\n"
        f"- Balanco do Mes: O saldo livre apos abater as despesas foi de R$ {saldo_livre:,.2f}. "
    )
    if economias_mes > 0:
        texto += f"Desse saldo livre, R$ {economias_mes:,.2f} foram transferidos e salvos em contas poupanca/guardados.\n"
    else:
        texto += "Nao foram registradas transferencias de economias para contas guardadas neste periodo.\n"
        
    texto += (
        f"\nEm relacao a terceiros (amigos/familiares):\n"
        f"- Cartao de Credito: Ha R$ {total_cartao_terceiros:,.2f} pendentes de reembolso em compras passadas no seu cartao para este mes.\n"
        f"- Emprestimos Diretos: O montante pendente de reembolso acumulado por emprestimos diretos (PIX/dinheiro) e de R$ {total_emprestimos_pendentes:,.2f}."
    )
    
    pdf.multi_cell(0, 5.5, texto.encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(8)
    
    # 2. Tabela de Indicadores
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "2. TABELA DE RESUMO MACRO DO MES", ln=True)
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    pdf.cell(100, 8, "Categoria de Fluxo", 1, 0, "L", fill=True)
    pdf.cell(80, 8, "Valor (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    indicadores = [
        ("Minhas Receitas (+)", total_receitas),
        ("Gastos Fixos (-)", total_fixos),
        ("Gastos Avulsos (-)", total_avulsos),
        ("Saldo Livre", saldo_livre),
        ("Economias Guardadas", economias_mes),
        ("Fatura de Terceiros a Receber (Mes)", total_cartao_terceiros),
        ("Emprestimos a Receber (Geral)", total_emprestimos_pendentes)
    ]
    for ind, val in indicadores:
        pdf.cell(100, 8, ind.encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(80, 8, f"R$ {val:,.2f}", 1, 1, "R")
        
    pdf.ln(8)
    
    # 3. Contas e Saldos
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "3. INTEGRACAO DE CONTAS, SALDOS E CARTOES", ln=True)
    pdf.set_font("Helvetica", "", 10.5)
    pdf.multi_cell(0, 5.5, "Contas bancarias com seus saldos de fechamento atuais e faturas de cartoes de credito pendentes:".encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(2)
    
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    pdf.cell(100, 8, "Conta / Banco / Cartao", 1, 0, "L", fill=True)
    pdf.cell(40, 8, "Tipo", 1, 0, "C", fill=True)
    pdf.cell(40, 8, "Saldo / Fatura (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    for conta in dados.get("contas", []):
        pdf.cell(100, 8, conta["nome"].encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(40, 8, conta.get("tipo", "Normal"), 1, 0, "C")
        pdf.cell(40, 8, f"R$ {conta['saldo']:,.2f}", 1, 1, "R")
        
    for cartao in dados.get("cartoes", []):
        pdf.cell(100, 8, (f"Cartao {cartao['nome']}").encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(40, 8, "Cartao", 1, 0, "C")
        pdf.cell(40, 8, f"R$ {cartao['fatura']:,.2f}", 1, 1, "R")
        
    # 4. Saldos a Receber de Terceiros
    pdf.ln(8)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "4. SALDOS A RECEBER DE TERCEIROS (DEVEDORES)", ln=True)
    
    nomes_devedores = sorted(list(set([t["nome"] for t in unpaid_cartao] + [e["nome"] for e in unpaid_emprestimo])))
    
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    pdf.cell(100, 8, "Pessoa / Devedor", 1, 0, "L", fill=True)
    pdf.cell(80, 8, "Total Devido (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    if not nomes_devedores:
        pdf.cell(180, 8, "Nao ha saldos pendentes de terceiros neste periodo.", 1, 1, "C")
    else:
        for nome_dev in nomes_devedores:
            total_dev = sum(t["valor"] for t in unpaid_cartao if t["nome"].lower() == nome_dev.lower()) + sum(e["valor"] for e in unpaid_emprestimo if e["nome"].lower() == nome_dev.lower())
            pdf.cell(100, 8, nome_dev.encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
            pdf.cell(80, 8, f"R$ {total_dev:,.2f}", 1, 1, "R")
            
    return bytes(pdf.output())

def gerar_relatorio_mensal_docx(dados, periodo, mes_nome, ano, usuario_atual):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    import io
    
    # Filtrar dados do mês selecionado
    receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo)]
    gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if pertence_periodo_cartao(g, periodo)]
    gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if pertence_periodo_cartao(g, periodo)]
    
    total_receitas = sum(r["valor"] for r in receitas_mes)
    total_fixos = sum(g["valor"] for g in gastos_fixos_mes)
    total_avulsos = sum(g["valor"] for g in gastos_avulsos_mes)
    total_gastos = total_fixos + total_avulsos
    saldo_livre = total_receitas - total_gastos
    
    # Economias
    economias_mes = 0.0
    for t in dados.get("transferencias", []):
        if t["data"].startswith(periodo):
            dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
            orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
            if dest_tipo == "Guardado" and orig_tipo == "Normal":
                economias_mes += t["valor"]
                
    # Terceiros pendentes
    unpaid_cartao = [t for t in dados.get("gastos_terceiros_cartao", []) if not t["pago"]]
    unpaid_emprestimo = [e for e in dados.get("gastos_terceiros_emprestimo", []) if not e["pago"]]
    
    total_cartao_terceiros = sum(t["valor"] for t in unpaid_cartao if pertence_periodo_cartao(t, periodo))
    total_emprestimos_pendentes = sum(e["valor"] for e in dados.get("gastos_terceiros_emprestimo", []) if emprestimo_pendente_no_periodo(e, periodo))
    
    doc = docx.Document()
    
    # Título do Relatório
    title = doc.add_paragraph()
    r_title = title.add_run(f"Relatório Financeiro Mensal - {mes_nome} / {ano}")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(130, 87, 229)
    
    subtitle = doc.add_paragraph()
    r_sub = subtitle.add_run(f"Finanças Didáticas | Usuário: {usuario_atual}\nGerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    
    doc.add_heading("1. Resumo Explicativo do Período", level=1)
    p1 = doc.add_paragraph()
    p1.add_run(f"Durante o mês de {mes_nome} de {ano}, o perfil '{usuario_atual}' realizou a organização financeira. A seguir, a análise do período:\n\n")
    p1.add_run("Receitas Registradas: ").bold = True
    p1.add_run(f"Entrou um montante de R$ {total_receitas:,.2f}.\n")
    p1.add_run("Despesas do Mês: ").bold = True
    p1.add_run(f"As despesas totalizaram R$ {total_gastos:,.2f}, compostas por R$ {total_fixos:,.2f} de custos fixos e R$ {total_avulsos:,.2f} de gastos avulsos.\n")
    p1.add_run("Saldo Livre: ").bold = True
    p1.add_run(f"O saldo livre restante foi de R$ {saldo_livre:,.2f}. ")
    if economias_mes > 0:
        p1.add_run(f"Destinou R$ {economias_mes:,.2f} para poupança.\n")
    else:
        p1.add_run("Não foram registrados lançamentos de economias direcionadas para a poupança neste mês.\n")
        
    p1.add_run("\nSobre Finanças de Terceiros:\n").bold = True
    p1.add_run(f"- Cartão de Crédito: Há R$ {total_cartao_terceiros:,.2f} pendentes de reembolso neste mês.\n")
    p1.add_run(f"- Empréstimos: O total pendente de devolução a você é de R$ {total_emprestimos_pendentes:,.2f}.\n")
    
    doc.add_heading("2. Tabela de Indicadores Financeiros", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Indicador'
    hdr_cells[0].paragraphs[0].runs[0].font.bold = True
    hdr_cells[1].text = 'Valor (R$)'
    hdr_cells[1].paragraphs[0].runs[0].font.bold = True
    
    indicadores = [
        ("Minhas Receitas (+)", total_receitas),
        ("Gastos Fixos (-)", total_fixos),
        ("Gastos Avulsos (-)", total_avulsos),
        ("Saldo Livre", saldo_livre),
        ("Economias Guardadas", economias_mes),
        ("Fatura de Terceiros a Receber (Mes)", total_cartao_terceiros),
        ("Empréstimos a Receber (Geral)", total_emprestimos_pendentes)
    ]
    for ind, val in indicadores:
        row_cells = table.add_row().cells
        row_cells[0].text = ind
        row_cells[1].text = f"R$ {val:,.2f}"
        
    for row in table.rows:
        row.cells[0].width = Inches(4.5)
        row.cells[1].width = Inches(2.0)
        
    doc.add_paragraph("\n")
    doc.add_heading("3. Fechamento das Contas, Saldos e Cartões", level=1)
    doc.add_paragraph("Abaixo estão listados os bancos e faturas com seus respectivos valores atuais:")
    
    table_contas = doc.add_table(rows=1, cols=3)
    table_contas.style = 'Table Grid'
    table_contas.autofit = False
    
    hdr_contas = table_contas.rows[0].cells
    hdr_contas[0].text = 'Conta / Banco / Cartão'
    hdr_contas[0].paragraphs[0].runs[0].font.bold = True
    hdr_contas[1].text = 'Tipo'
    hdr_contas[1].paragraphs[0].runs[0].font.bold = True
    hdr_contas[2].text = 'Saldo / Fatura (R$)'
    hdr_contas[2].paragraphs[0].runs[0].font.bold = True
    
    for conta in dados.get("contas", []):
        row_cells = table_contas.add_row().cells
        row_cells[0].text = conta["nome"]
        row_cells[1].text = conta.get("tipo", "Normal")
        row_cells[2].text = f"R$ {conta['saldo']:,.2f}"
        
    for cartao in dados.get("cartoes", []):
        row_cells = table_contas.add_row().cells
        row_cells[0].text = f"Cartão {cartao['nome']}"
        row_cells[1].text = "Cartão de Crédito"
        row_cells[2].text = f"R$ {cartao['fatura']:,.2f}"
        
    for row in table_contas.rows:
        row.cells[0].width = Inches(3.0)
        row.cells[1].width = Inches(1.5)
        row.cells[2].width = Inches(2.0)
        
    # 4. Saldos a receber de terceiros
    doc.add_paragraph("\n")
    doc.add_heading("4. Saldos Pendentes por Devedor (Terceiros)", level=1)
    
    nomes_devedores = sorted(list(set([t["nome"] for t in unpaid_cartao] + [e["nome"] for e in unpaid_emprestimo])))
    
    table_dev = doc.add_table(rows=1, cols=2)
    table_dev.style = 'Table Grid'
    table_dev.autofit = False
    
    hdr_dev = table_dev.rows[0].cells
    hdr_dev[0].text = 'Nome do Devedor'
    hdr_dev[0].paragraphs[0].runs[0].font.bold = True
    hdr_dev[1].text = 'Valor Total Devido (R$)'
    hdr_dev[1].paragraphs[0].runs[0].font.bold = True
    
    if not nomes_devedores:
        row_cells = table_dev.add_row().cells
        row_cells[0].text = "Sem pendências"
        row_cells[1].text = "R$ 0,00"
    else:
        for nome_dev in nomes_devedores:
            total_dev = sum(t["valor"] for t in unpaid_cartao if t["nome"].lower() == nome_dev.lower()) + sum(e["valor"] for e in unpaid_emprestimo if e["nome"].lower() == nome_dev.lower())
            row_cells = table_dev.add_row().cells
            row_cells[0].text = nome_dev
            row_cells[1].text = f"R$ {total_dev:,.2f}"
            
    for row in table_dev.rows:
        row.cells[0].width = Inches(4.5)
        row.cells[1].width = Inches(2.0)
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def gerar_relatorio_anual_pdf(dados, ano, meses_lista, usuario_atual):
    from fpdf import FPDF
    from datetime import datetime
    
    dados_ano = []
    for m in range(1, 13):
        prefixo_busca = f"{ano}-{m:02d}"
        rec_ano = sum(r["valor"] for r in dados.get("receitas", []) if r["data"].startswith(prefixo_busca))
        fix_ano = sum(g["valor"] for g in dados.get("gastos_fixos", []) if g["data"].startswith(prefixo_busca))
        av_ano = sum(g["valor"] for g in dados.get("gastos_avulsos", []) if g["data"].startswith(prefixo_busca))
        gastos_totais = fix_ano + av_ano
        
        # Calcular economias
        economias_mes = 0.0
        for t in dados.get("transferencias", []):
            if t["data"].startswith(prefixo_busca):
                dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
                orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
                if dest_tipo == "Guardado" and orig_tipo == "Normal":
                    economias_mes += t["valor"]
                    
        dados_ano.append({
            "mes_num": m,
            "mes_nome": meses_lista[m-1],
            "ganhos": rec_ano,
            "gastos": gastos_totais,
            "guardado": economias_mes,
            "balanco": rec_ano - gastos_totais
        })
        
    total_ganhos_ano = sum(d["ganhos"] for d in dados_ano)
    total_gastos_ano = sum(d["gastos"] for d in dados_ano)
    total_guardado_ano = sum(d["guardado"] for d in dados_ano)
    saldo_anual_acumulado = total_ganhos_ano - total_gastos_ano
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_margins(15, 15, 15)
    pdf.set_text_color(32, 32, 36)
    
    # Cabeçalho
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "FINANCAS DIDATICAS - RELATORIO ANUAL CONSOLIDADO", ln=True, align="C")
    pdf.set_font("Helvetica", "", 12)
    pdf.cell(0, 8, f"Ano de Referencia: {ano}", ln=True, align="C")
    pdf.cell(0, 6, f"Usuario: {usuario_atual} | Gerado em: {datetime.now().strftime('%d/%m/%Y as %H:%M')}", ln=True, align="C")
    pdf.ln(8)
    
    # Linha divisória Roxo
    pdf.set_draw_color(130, 87, 229)
    pdf.set_line_width(1)
    pdf.line(15, pdf.get_y(), 195, pdf.get_y())
    pdf.ln(6)
    
    # 1. Resumo Explicativo do Ano
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "1. RESUMO EXPLICATIVO DO ANO CONSOLIDADO", ln=True)
    pdf.set_font("Helvetica", "", 10.5)
    
    texto = (
        f"Este relatorio anual consolida a saude financeira do perfil '{usuario_atual}' ao longo de todo o ano de {ano}. "
        f"Abaixo estao as explicacoes e as metricas acumuladas do periodo:\n\n"
        f"- Ganhos Acumulados: Ao longo de todo o ano de {ano}, voce registrou um faturamento bruto acumulado de R$ {total_ganhos_ano:,.2f}.\n"
        f"- Despesas Acumuladas: O total de despesas reais somou R$ {total_gastos_ano:,.2f} (englobando todas as contas fixas, parcelamentos e gastos avulsos).\n"
        f"- Balanco Anual: O balanco financeiro liquido acumulado foi de R$ {saldo_anual_acumulado:,.2f}. "
    )
    if total_guardado_ano > 0:
        texto += f"Durante o ano, voce conseguiu guardar com sucesso um total de R$ {total_guardado_ano:,.2f} em economias direcionadas para contas poupancas."
    else:
        texto += "Nao foram registradas transferencias de economias para poupancas de longo prazo durante o ano."
        
    pdf.multi_cell(0, 5.5, texto.encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(8)
    
    # 2. Tabela de Consolidação Mensal
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "2. TABELA DE COMPOSICAO MES A MES", ln=True)
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    
    pdf.cell(50, 8, "Mes", 1, 0, "L", fill=True)
    pdf.cell(40, 8, "Ganhos (R$)", 1, 0, "R", fill=True)
    pdf.cell(40, 8, "Gastos (R$)", 1, 0, "R", fill=True)
    pdf.cell(50, 8, "Balanco Liquido (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    for d in dados_ano:
        pdf.cell(50, 8, d["mes_nome"].encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(40, 8, f"R$ {d['ganhos']:,.2f}", 1, 0, "R")
        pdf.cell(40, 8, f"R$ {d['gastos']:,.2f}", 1, 0, "R")
        pdf.cell(50, 8, f"R$ {d['balanco']:,.2f}", 1, 1, "R")
        
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(230, 230, 235)
    pdf.cell(50, 8, "TOTAL ANUAL", 1, 0, "L", fill=True)
    pdf.cell(40, 8, f"R$ {total_ganhos_ano:,.2f}", 1, 0, "R", fill=True)
    pdf.cell(40, 8, f"R$ {total_gastos_ano:,.2f}", 1, 0, "R", fill=True)
    pdf.cell(50, 8, f"R$ {saldo_anual_acumulado:,.2f}", 1, 1, "R", fill=True)
    
    return bytes(pdf.output())

def gerar_relatorio_anual_docx(dados, ano, meses_lista, usuario_atual):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    import io
    from datetime import datetime
    
    dados_ano = []
    for m in range(1, 13):
        prefixo_busca = f"{ano}-{m:02d}"
        rec_ano = sum(r["valor"] for r in dados.get("receitas", []) if r["data"].startswith(prefixo_busca))
        fix_ano = sum(g["valor"] for g in dados.get("gastos_fixos", []) if g["data"].startswith(prefixo_busca))
        av_ano = sum(g["valor"] for g in dados.get("gastos_avulsos", []) if g["data"].startswith(prefixo_busca))
        gastos_totais = fix_ano + av_ano
        
        # Calcular economias
        economias_mes = 0.0
        for t in dados.get("transferencias", []):
            if t["data"].startswith(prefixo_busca):
                dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
                orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
                if dest_tipo == "Guardado" and orig_tipo == "Normal":
                    economias_mes += t["valor"]
                    
        dados_ano.append({
            "mes_nome": meses_lista[m-1],
            "ganhos": rec_ano,
            "gastos": gastos_totais,
            "guardado": economias_mes,
            "balanco": rec_ano - gastos_totais
        })
        
    total_ganhos_ano = sum(d["ganhos"] for d in dados_ano)
    total_gastos_ano = sum(d["gastos"] for d in dados_ano)
    total_guardado_ano = sum(d["guardado"] for d in dados_ano)
    saldo_anual_acumulado = total_ganhos_ano - total_gastos_ano
    
    doc = docx.Document()
    
    # Título do Relatório
    title = doc.add_paragraph()
    r_title = title.add_run(f"Relatório Financeiro Anual Consolidado - {ano}")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(130, 87, 229)
    
    subtitle = doc.add_paragraph()
    r_sub = subtitle.add_run(f"Finanças Didáticas | Usuário: {usuario_atual}\nGerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    
    doc.add_heading("1. Resumo Explicativo do Ano", level=1)
    p1 = doc.add_paragraph()
    p1.add_run(f"O presente relatório consolida o desempenho financeiro do perfil '{usuario_atual}' ao longo do ano de {ano}. A seguir, apresenta-se a síntese geral das movimentações:\n\n")
    p1.add_run("Ganhos Anuais Acumulados: ").bold = True
    p1.add_run(f"O faturamento bruto do ano somou R$ {total_ganhos_ano:,.2f}.\n")
    p1.add_run("Despesas Anuais Acumuladas: ").bold = True
    p1.add_run(f"O montante de saídas reais somou R$ {total_gastos_ano:,.2f}.\n")
    p1.add_run("Balanço Anual: ").bold = True
    p1.add_run(f"O balanço financeiro líquido acumulado foi de R$ {saldo_anual_acumulado:,.2f}. ")
    if total_guardado_ano > 0:
        p1.add_run(f"Deste valor, R$ {total_guardado_ano:,.2f} foram economizados com sucesso nas contas de poupança/investimentos (Guardados).\n")
    else:
        p1.add_run("Não foram registradas transferências de economias direcionadas para a poupança neste ano.\n")
        
    doc.add_heading("2. Detalhamento Mensal (Mês a Mês)", level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Mês'
    hdr_cells[0].paragraphs[0].runs[0].font.bold = True
    hdr_cells[1].text = 'Ganhos (R$)'
    hdr_cells[1].paragraphs[0].runs[0].font.bold = True
    hdr_cells[2].text = 'Gastos (R$)'
    hdr_cells[2].paragraphs[0].runs[0].font.bold = True
    hdr_cells[3].text = 'Balanço Líquido (R$)'
    hdr_cells[3].paragraphs[0].runs[0].font.bold = True
    
    for d in dados_ano:
        row_cells = table.add_row().cells
        row_cells[0].text = d["mes_nome"]
        row_cells[1].text = f"R$ {d['ganhos']:,.2f}"
        row_cells[2].text = f"R$ {d['gastos']:,.2f}"
        row_cells[3].text = f"R$ {d['balanco']:,.2f}"
        
    # Linha final de total
    total_row = table.add_row().cells
    total_row[0].text = "TOTAL ANUAL"
    total_row[0].paragraphs[0].runs[0].font.bold = True
    total_row[1].text = f"R$ {total_ganhos_ano:,.2f}"
    total_row[1].paragraphs[0].runs[0].font.bold = True
    total_row[2].text = f"R$ {total_gastos_ano:,.2f}"
    total_row[2].paragraphs[0].runs[0].font.bold = True
    total_row[3].text = f"R$ {saldo_anual_acumulado:,.2f}"
    total_row[3].paragraphs[0].runs[0].font.bold = True
    
    for row in table.rows:
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(1.5)
        row.cells[2].width = Inches(1.5)
        row.cells[3].width = Inches(1.5)
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
