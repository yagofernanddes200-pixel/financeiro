import streamlit as st
import pandas as pd
import plotly.express as px
import json
import os
import re
import hashlib
from datetime import datetime, timedelta

# Configuração da página para celular (Samsung A15 e iPhone)
st.set_page_config(
    page_title="Finanças Didáticas",
    page_icon="💰",
    layout="centered",
    initial_sidebar_state="collapsed"
)

# --- CONFIGURAÇÃO DE ESTILO: MODO ESCURO PREMIUM ---
st.markdown("""
<style>
    /* Estilo geral da página */
    .stApp {
        background-color: #121214;
        color: #E1E1E6;
    }
    
    /* Cabeçalho */
    h1, h2, h3, h4, h5, h6 {
        color: #FFFFFF !important;
        font-family: 'Inter', sans-serif;
    }
    
    /* Cards Customizados */
    .card-principal {
        background-color: #202024;
        padding: 18px;
        border-radius: 12px;
        margin-bottom: 15px;
        border-left: 5px solid #8257E5;
        box-shadow: 0 4px 6px rgba(0, 0, 0, 0.3);
    }
    .card-didatico {
        background-color: #1A1A1E;
        padding: 15px;
        border-radius: 10px;
        margin-bottom: 15px;
        border-left: 5px solid #3867D6;
        font-size: 0.9em;
        line-height: 1.5;
        color: #C4C4CC;
    }
    .card-receita {
        background-color: #1A2E26;
        padding: 15px;
        border-radius: 10px;
        border-left: 5px solid #26DE81;
    }
    .card-gasto {
        background-color: #2D1A1E;
        padding: 15px;
        border-radius: 10px;
        border-left: 5px solid #FF4757;
    }
    .titulo-secao {
        color: #FFFFFF;
        font-weight: bold;
        margin-top: 20px;
        border-bottom: 1px solid #29292E;
        padding-bottom: 8px;
    }
    
    /* Botões */
    .stButton>button {
        background-color: #8257E5 !important;
        color: white !important;
        border-radius: 8px !important;
        border: none !important;
        padding: 10px 20px !important;
        font-weight: bold !important;
        width: 100%;
        transition: background-color 0.2s;
    }
    .stButton>button:hover {
        background-color: #9466FF !important;
    }
    .stButton>button:disabled {
        background-color: #3A3A40 !important;
        color: #7A7A82 !important;
        cursor: not-allowed;
    }
    
    /* Inputs */
    div[data-baseweb="input"] {
        background-color: #202024 !important;
        border-color: #29292E !important;
    }
    input {
        color: #FFFFFF !important;
    }
</style>
""", unsafe_allow_html=True)

# --- FUNÇÕES DE AUXÍLIO PARA DATAS ---
def get_proximos_meses(data_inicial, parcelas):
    """Retorna uma lista de strings 'AAAA-MM' para as parcelas futuras."""
    datas = []
    ano = data_inicial.year
    mes = data_inicial.month
    
    for i in range(parcelas):
        datas.append(f"{ano}-{mes:02d}")
        mes += 1
        if mes > 12:
            mes = 1
            ano += 1
    return datas


def formatar_data_br(data_iso):
    """Converte 'AAAA-MM-DD' para 'DD/MM/AAAA'. Retorna o original se não conseguir converter."""
    try:
        return datetime.strptime(str(data_iso), "%Y-%m-%d").strftime("%d/%m/%Y")
    except (ValueError, TypeError):
        return str(data_iso)


def periodo_anterior(periodo):
    """Recebe 'AAAA-MM' e retorna o período do mês anterior."""
    ano, mes = map(int, periodo.split("-"))
    mes -= 1
    if mes < 1:
        mes = 12
        ano -= 1
    return f"{ano}-{mes:02d}"


def periodo_seguinte(periodo):
    """Recebe 'AAAA-MM' e retorna o período do mês seguinte."""
    ano, mes = map(int, periodo.split("-"))
    mes += 1
    if mes > 12:
        mes = 1
        ano += 1
    return f"{ano}-{mes:02d}"


def calcular_periodo_fatura(data_iso, dia_fechamento):
    """
    Calcula em qual fatura (AAAA-MM) uma compra no cartão cai, considerando o dia de fechamento.
    Se a compra ocorreu NO dia do fechamento ou depois, ela entra na fatura do mês seguinte.
    """
    try:
        data_compra = datetime.strptime(str(data_iso), "%Y-%m-%d")
    except (ValueError, TypeError):
        return str(data_iso)[:7]

    periodo_atual = f"{data_compra.year}-{data_compra.month:02d}"
    if data_compra.day >= int(dia_fechamento):
        return periodo_seguinte(periodo_atual)
    return periodo_atual


def obter_dados_cartao(dados, nome_cartao):
    return next((c for c in dados.get("cartoes", []) if c["nome"] == nome_cartao), None)


def pertence_periodo_cartao(item, periodo):
    """
    Verifica se um lançamento pertence a um período (AAAA-MM).
    Usa o campo 'periodo_fatura' (calculado a partir do fechamento do cartão) quando existir;
    caso contrário, cai de volta para a data literal do lançamento.
    """
    return item.get("periodo_fatura", str(item["data"])[:7]) == periodo


def hash_resposta(resposta):
    return hashlib.sha256(resposta.strip().lower().encode("utf-8")).hexdigest()


PERGUNTAS_SEGURANCA_PADRAO = [
    "Qual o nome do seu primeiro animal de estimação?",
    "Qual o nome da cidade onde você nasceu?",
    "Qual o nome dos seus filhos?",
    "Qual o apelido que você tinha na infância?",
    "Personalizada (escrever a minha própria pergunta)"
]


# --- FUNÇÕES PARA GERAÇÃO DE RELATÓRIOS (PDF & WORD) ---
def gerar_relatorio_mensal_pdf(dados, periodo, mes_nome, ano, usuario_atual):
    from fpdf import FPDF
    from datetime import datetime
    
    # Filtrar dados do mês selecionado
    receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo)]
    gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if pertence_periodo_cartao(g, periodo)]
    gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if pertence_periodo_cartao(g, periodo)]
    
    total_receitas = sum(r["valor"] for r in receitas_mes)
    total_fixos = sum(g["valor"] for g in gastos_fixos_mes)
    total_avulsos = sum(g["valor"] for g in gastos_avulsos_mes)
    total_gastos = total_fixos + total_avulsos
    saldo_livre = total_receitas - total_gastos
    
    # Economias do mês
    economias_mes = 0.0
    for t in dados.get("transferencias", []):
        if t["data"].startswith(periodo):
            dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
            orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
            if dest_tipo == "Guardado" and orig_tipo == "Normal":
                economias_mes += t["valor"]
                
    # Terceiros pendentes
    unpaid_cartao = [t for t in dados.get("gastos_terceiros_cartao", []) if not t["pago"]]
    unpaid_emprestimo = [e for e in dados.get("gastos_terceiros_emprestimo", []) if not e["pago"]]
    
    total_cartao_terceiros = sum(t["valor"] for t in unpaid_cartao if pertence_periodo_cartao(t, periodo))
    total_emprestimos_pendentes = sum(e["valor"] for e in unpaid_emprestimo)
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_margins(15, 15, 15)
    pdf.set_text_color(32, 32, 36)
    
    # Cabeçalho
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "FINANCAS DIDATICAS - RELATORIO MENSAL", ln=True, align="C")
    pdf.set_font("Helvetica", "", 12)
    pdf.cell(0, 8, f"Mes de Referencia: {mes_nome} de {ano}", ln=True, align="C")
    pdf.cell(0, 6, f"Usuario: {usuario_atual} | Gerado em: {datetime.now().strftime('%d/%m/%Y as %H:%M')}", ln=True, align="C")
    pdf.ln(8)
    
    # Linha divisória Roxo
    pdf.set_draw_color(130, 87, 229)
    pdf.set_line_width(1)
    pdf.line(15, pdf.get_y(), 195, pdf.get_y())
    pdf.ln(6)
    
    # 1. Resumo Explicativo
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "1. RESUMO EXPLICATIVO DO PERIODO", ln=True)
    pdf.set_font("Helvetica", "", 10.5)
    
    texto = (
        f"Durante o mes de {mes_nome} de {ano}, o perfil '{usuario_atual}' apresentou as seguintes movimentacoes:\n\n"
        f"- Receitas Totais: Entrou um total de R$ {total_receitas:,.2f} em receitas cadastradas.\n"
        f"- Custos Totais: O total de despesas reais pagas ou a pagar no mes somou R$ {total_gastos:,.2f}, dividindo-se em R$ {total_fixos:,.2f} de gastos fixos/parcelamentos e R$ {total_avulsos:,.2f} de gastos avulsos (dia a dia).\n"
        f"- Balanco do Mes: O saldo livre apos abater as despesas foi de R$ {saldo_livre:,.2f}. "
    )
    if economias_mes > 0:
        texto += f"Desse saldo livre, R$ {economias_mes:,.2f} foram transferidos e salvos em contas poupanca/guardados.\n"
    else:
        texto += "Nao foram registradas transferencias de economias para contas guardadas neste periodo.\n"
        
    texto += (
        f"\nEm relacao a terceiros (amigos/familiares):\n"
        f"- Cartao de Credito: Ha R$ {total_cartao_terceiros:,.2f} pendentes de reembolso em compras passadas no seu cartao para este mes.\n"
        f"- Emprestimos Diretos: O montante pendente de reembolso acumulado por emprestimos diretos (PIX/dinheiro) e de R$ {total_emprestimos_pendentes:,.2f}."
    )
    
    pdf.multi_cell(0, 5.5, texto.encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(8)
    
    # 2. Tabela de Indicadores
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "2. TABELA DE RESUMO MACRO DO MES", ln=True)
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    pdf.cell(100, 8, "Categoria de Fluxo", 1, 0, "L", fill=True)
    pdf.cell(80, 8, "Valor (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    indicadores = [
        ("Minhas Receitas (+)", total_receitas),
        ("Gastos Fixos (-)", total_fixos),
        ("Gastos Avulsos (-)", total_avulsos),
        ("Saldo Livre", saldo_livre),
        ("Economias Guardadas", economias_mes),
        ("Fatura de Terceiros a Receber (Mes)", total_cartao_terceiros),
        ("Emprestimos a Receber (Geral)", total_emprestimos_pendentes)
    ]
    for ind, val in indicadores:
        pdf.cell(100, 8, ind.encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(80, 8, f"R$ {val:,.2f}", 1, 1, "R")
        
    pdf.ln(8)
    
    # 3. Contas e Saldos
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "3. INTEGRACAO DE CONTAS, SALDOS E CARTOES", ln=True)
    pdf.set_font("Helvetica", "", 10.5)
    pdf.multi_cell(0, 5.5, "Contas bancarias com seus saldos de fechamento atuais e faturas de cartoes de credito pendentes:".encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(2)
    
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    pdf.cell(100, 8, "Conta / Banco / Cartao", 1, 0, "L", fill=True)
    pdf.cell(40, 8, "Tipo", 1, 0, "C", fill=True)
    pdf.cell(40, 8, "Saldo / Fatura (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    for conta in dados.get("contas", []):
        pdf.cell(100, 8, conta["nome"].encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(40, 8, conta.get("tipo", "Normal"), 1, 0, "C")
        pdf.cell(40, 8, f"R$ {conta['saldo']:,.2f}", 1, 1, "R")
        
    for cartao in dados.get("cartoes", []):
        pdf.cell(100, 8, (f"Cartao {cartao['nome']}").encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(40, 8, "Cartao", 1, 0, "C")
        pdf.cell(40, 8, f"R$ {cartao['fatura']:,.2f}", 1, 1, "R")
        
    # 4. Saldos a Receber de Terceiros
    pdf.ln(8)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "4. SALDOS A RECEBER DE TERCEIROS (DEVEDORES)", ln=True)
    
    nomes_devedores = sorted(list(set([t["nome"] for t in unpaid_cartao] + [e["nome"] for e in unpaid_emprestimo])))
    
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    pdf.cell(100, 8, "Pessoa / Devedor", 1, 0, "L", fill=True)
    pdf.cell(80, 8, "Total Devido (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    if not nomes_devedores:
        pdf.cell(180, 8, "Nao ha saldos pendentes de terceiros neste periodo.", 1, 1, "C")
    else:
        for nome_dev in nomes_devedores:
            total_dev = sum(t["valor"] for t in unpaid_cartao if t["nome"].lower() == nome_dev.lower()) + sum(e["valor"] for e in unpaid_emprestimo if e["nome"].lower() == nome_dev.lower())
            pdf.cell(100, 8, nome_dev.encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
            pdf.cell(80, 8, f"R$ {total_dev:,.2f}", 1, 1, "R")
            
    return bytes(pdf.output())

def gerar_relatorio_mensal_docx(dados, periodo, mes_nome, ano, usuario_atual):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    import io
    
    # Filtrar dados do mês selecionado
    receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo)]
    gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if pertence_periodo_cartao(g, periodo)]
    gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if pertence_periodo_cartao(g, periodo)]
    
    total_receitas = sum(r["valor"] for r in receitas_mes)
    total_fixos = sum(g["valor"] for g in gastos_fixos_mes)
    total_avulsos = sum(g["valor"] for g in gastos_avulsos_mes)
    total_gastos = total_fixos + total_avulsos
    saldo_livre = total_receitas - total_gastos
    
    # Economias
    economias_mes = 0.0
    for t in dados.get("transferencias", []):
        if t["data"].startswith(periodo):
            dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
            orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
            if dest_tipo == "Guardado" and orig_tipo == "Normal":
                economias_mes += t["valor"]
                
    # Terceiros pendentes
    unpaid_cartao = [t for t in dados.get("gastos_terceiros_cartao", []) if not t["pago"]]
    unpaid_emprestimo = [e for e in dados.get("gastos_terceiros_emprestimo", []) if not e["pago"]]
    
    total_cartao_terceiros = sum(t["valor"] for t in unpaid_cartao if pertence_periodo_cartao(t, periodo))
    total_emprestimos_pendentes = sum(e["valor"] for e in unpaid_emprestimo)
    
    doc = docx.Document()
    
    # Título do Relatório
    title = doc.add_paragraph()
    r_title = title.add_run(f"Relatório Financeiro Mensal - {mes_nome} / {ano}")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(130, 87, 229)
    
    subtitle = doc.add_paragraph()
    r_sub = subtitle.add_run(f"Finanças Didáticas | Usuário: {usuario_atual}\nGerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    
    doc.add_heading("1. Resumo Explicativo do Período", level=1)
    p1 = doc.add_paragraph()
    p1.add_run(f"Durante o mês de {mes_nome} de {ano}, o perfil '{usuario_atual}' realizou a organização financeira. A seguir, a análise do período:\n\n")
    p1.add_run("Receitas Registradas: ").bold = True
    p1.add_run(f"Entrou um montante de R$ {total_receitas:,.2f}.\n")
    p1.add_run("Despesas do Mês: ").bold = True
    p1.add_run(f"As despesas totalizaram R$ {total_gastos:,.2f}, compostas por R$ {total_fixos:,.2f} de custos fixos e R$ {total_avulsos:,.2f} de gastos avulsos.\n")
    p1.add_run("Saldo Livre: ").bold = True
    p1.add_run(f"O saldo livre restante foi de R$ {saldo_livre:,.2f}. ")
    if economias_mes > 0:
        p1.add_run(f"Destinou R$ {economias_mes:,.2f} para poupança.\n")
    else:
        p1.add_run("Não foram registrados lançamentos de economias direcionadas para a poupança neste mês.\n")
        
    p1.add_run("\nSobre Finanças de Terceiros:\n").bold = True
    p1.add_run(f"- Cartão de Crédito: Há R$ {total_cartao_terceiros:,.2f} pendentes de reembolso neste mês.\n")
    p1.add_run(f"- Empréstimos: O total pendente de devolução a você é de R$ {total_emprestimos_pendentes:,.2f}.\n")
    
    doc.add_heading("2. Tabela de Indicadores Financeiros", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Indicador'
    hdr_cells[0].paragraphs[0].runs[0].font.bold = True
    hdr_cells[1].text = 'Valor (R$)'
    hdr_cells[1].paragraphs[0].runs[0].font.bold = True
    
    indicadores = [
        ("Minhas Receitas (+)", total_receitas),
        ("Gastos Fixos (-)", total_fixos),
        ("Gastos Avulsos (-)", total_avulsos),
        ("Saldo Livre", saldo_livre),
        ("Economias Guardadas", economias_mes),
        ("Fatura de Terceiros a Receber (Mes)", total_cartao_terceiros),
        ("Empréstimos a Receber (Geral)", total_emprestimos_pendentes)
    ]
    for ind, val in indicadores:
        row_cells = table.add_row().cells
        row_cells[0].text = ind
        row_cells[1].text = f"R$ {val:,.2f}"
        
    for row in table.rows:
        row.cells[0].width = Inches(4.5)
        row.cells[1].width = Inches(2.0)
        
    doc.add_paragraph("\n")
    doc.add_heading("3. Fechamento das Contas, Saldos e Cartões", level=1)
    doc.add_paragraph("Abaixo estão listados os bancos e faturas com seus respectivos valores atuais:")
    
    table_contas = doc.add_table(rows=1, cols=3)
    table_contas.style = 'Table Grid'
    table_contas.autofit = False
    
    hdr_contas = table_contas.rows[0].cells
    hdr_contas[0].text = 'Conta / Banco / Cartão'
    hdr_contas[0].paragraphs[0].runs[0].font.bold = True
    hdr_contas[1].text = 'Tipo'
    hdr_contas[1].paragraphs[0].runs[0].font.bold = True
    hdr_contas[2].text = 'Saldo / Fatura (R$)'
    hdr_contas[2].paragraphs[0].runs[0].font.bold = True
    
    for conta in dados.get("contas", []):
        row_cells = table_contas.add_row().cells
        row_cells[0].text = conta["nome"]
        row_cells[1].text = conta.get("tipo", "Normal")
        row_cells[2].text = f"R$ {conta['saldo']:,.2f}"
        
    for cartao in dados.get("cartoes", []):
        row_cells = table_contas.add_row().cells
        row_cells[0].text = f"Cartão {cartao['nome']}"
        row_cells[1].text = "Cartão de Crédito"
        row_cells[2].text = f"R$ {cartao['fatura']:,.2f}"
        
    for row in table_contas.rows:
        row.cells[0].width = Inches(3.0)
        row.cells[1].width = Inches(1.5)
        row.cells[2].width = Inches(2.0)
        
    # 4. Saldos a receber de terceiros
    doc.add_paragraph("\n")
    doc.add_heading("4. Saldos Pendentes por Devedor (Terceiros)", level=1)
    
    nomes_devedores = sorted(list(set([t["nome"] for t in unpaid_cartao] + [e["nome"] for e in unpaid_emprestimo])))
    
    table_dev = doc.add_table(rows=1, cols=2)
    table_dev.style = 'Table Grid'
    table_dev.autofit = False
    
    hdr_dev = table_dev.rows[0].cells
    hdr_dev[0].text = 'Nome do Devedor'
    hdr_dev[0].paragraphs[0].runs[0].font.bold = True
    hdr_dev[1].text = 'Valor Total Devido (R$)'
    hdr_dev[1].paragraphs[0].runs[0].font.bold = True
    
    if not nomes_devedores:
        row_cells = table_dev.add_row().cells
        row_cells[0].text = "Sem pendências"
        row_cells[1].text = "R$ 0,00"
    else:
        for nome_dev in nomes_devedores:
            total_dev = sum(t["valor"] for t in unpaid_cartao if t["nome"].lower() == nome_dev.lower()) + sum(e["valor"] for e in unpaid_emprestimo if e["nome"].lower() == nome_dev.lower())
            row_cells = table_dev.add_row().cells
            row_cells[0].text = nome_dev
            row_cells[1].text = f"R$ {total_dev:,.2f}"
            
    for row in table_dev.rows:
        row.cells[0].width = Inches(4.5)
        row.cells[1].width = Inches(2.0)
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def gerar_relatorio_anual_pdf(dados, ano, meses_lista, usuario_atual):
    from fpdf import FPDF
    from datetime import datetime
    
    dados_ano = []
    for m in range(1, 13):
        prefixo_busca = f"{ano}-{m:02d}"
        rec_ano = sum(r["valor"] for r in dados.get("receitas", []) if r["data"].startswith(prefixo_busca))
        fix_ano = sum(g["valor"] for g in dados.get("gastos_fixos", []) if g["data"].startswith(prefixo_busca))
        av_ano = sum(g["valor"] for g in dados.get("gastos_avulsos", []) if g["data"].startswith(prefixo_busca))
        gastos_totais = fix_ano + av_ano
        
        # Calcular economias
        economias_mes = 0.0
        for t in dados.get("transferencias", []):
            if t["data"].startswith(prefixo_busca):
                dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
                orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
                if dest_tipo == "Guardado" and orig_tipo == "Normal":
                    economias_mes += t["valor"]
                    
        dados_ano.append({
            "mes_num": m,
            "mes_nome": meses_lista[m-1],
            "ganhos": rec_ano,
            "gastos": gastos_totais,
            "guardado": economias_mes,
            "balanco": rec_ano - gastos_totais
        })
        
    total_ganhos_ano = sum(d["ganhos"] for d in dados_ano)
    total_gastos_ano = sum(d["gastos"] for d in dados_ano)
    total_guardado_ano = sum(d["guardado"] for d in dados_ano)
    saldo_anual_acumulado = total_ganhos_ano - total_gastos_ano
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_margins(15, 15, 15)
    pdf.set_text_color(32, 32, 36)
    
    # Cabeçalho
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "FINANCAS DIDATICAS - RELATORIO ANUAL CONSOLIDADO", ln=True, align="C")
    pdf.set_font("Helvetica", "", 12)
    pdf.cell(0, 8, f"Ano de Referencia: {ano}", ln=True, align="C")
    pdf.cell(0, 6, f"Usuario: {usuario_atual} | Gerado em: {datetime.now().strftime('%d/%m/%Y as %H:%M')}", ln=True, align="C")
    pdf.ln(8)
    
    # Linha divisória Roxo
    pdf.set_draw_color(130, 87, 229)
    pdf.set_line_width(1)
    pdf.line(15, pdf.get_y(), 195, pdf.get_y())
    pdf.ln(6)
    
    # 1. Resumo Explicativo do Ano
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "1. RESUMO EXPLICATIVO DO ANO CONSOLIDADO", ln=True)
    pdf.set_font("Helvetica", "", 10.5)
    
    texto = (
        f"Este relatorio anual consolida a saude financeira do perfil '{usuario_atual}' ao longo de todo o ano de {ano}. "
        f"Abaixo estao as explicacoes e as metricas acumuladas do periodo:\n\n"
        f"- Ganhos Acumulados: Ao longo de todo o ano de {ano}, voce registrou um faturamento bruto acumulado de R$ {total_ganhos_ano:,.2f}.\n"
        f"- Despesas Acumuladas: O total de despesas reais somou R$ {total_gastos_ano:,.2f} (englobando todas as contas fixas, parcelamentos e gastos avulsos).\n"
        f"- Balanco Anual: O balanco financeiro liquido acumulado foi de R$ {saldo_anual_acumulado:,.2f}. "
    )
    if total_guardado_ano > 0:
        texto += f"Durante o ano, voce conseguiu guardar com sucesso um total de R$ {total_guardado_ano:,.2f} em economias direcionadas para contas poupancas."
    else:
        texto += "Nao foram registradas transferencias de economias para poupancas de longo prazo durante o ano."
        
    pdf.multi_cell(0, 5.5, texto.encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(8)
    
    # 2. Tabela de Consolidação Mensal
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "2. TABELA DE COMPOSICAO MES A MES", ln=True)
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(240, 242, 246)
    
    pdf.cell(50, 8, "Mes", 1, 0, "L", fill=True)
    pdf.cell(40, 8, "Ganhos (R$)", 1, 0, "R", fill=True)
    pdf.cell(40, 8, "Gastos (R$)", 1, 0, "R", fill=True)
    pdf.cell(50, 8, "Balanco Liquido (R$)", 1, 1, "R", fill=True)
    
    pdf.set_font("Helvetica", "", 10)
    for d in dados_ano:
        pdf.cell(50, 8, d["mes_nome"].encode('latin-1', 'replace').decode('latin-1'), 1, 0, "L")
        pdf.cell(40, 8, f"R$ {d['ganhos']:,.2f}", 1, 0, "R")
        pdf.cell(40, 8, f"R$ {d['gastos']:,.2f}", 1, 0, "R")
        pdf.cell(50, 8, f"R$ {d['balanco']:,.2f}", 1, 1, "R")
        
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(230, 230, 235)
    pdf.cell(50, 8, "TOTAL ANUAL", 1, 0, "L", fill=True)
    pdf.cell(40, 8, f"R$ {total_ganhos_ano:,.2f}", 1, 0, "R", fill=True)
    pdf.cell(40, 8, f"R$ {total_gastos_ano:,.2f}", 1, 0, "R", fill=True)
    pdf.cell(50, 8, f"R$ {saldo_anual_acumulado:,.2f}", 1, 1, "R", fill=True)
    
    return bytes(pdf.output())

def gerar_relatorio_anual_docx(dados, ano, meses_lista, usuario_atual):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    import io
    from datetime import datetime
    
    dados_ano = []
    for m in range(1, 13):
        prefixo_busca = f"{ano}-{m:02d}"
        rec_ano = sum(r["valor"] for r in dados.get("receitas", []) if r["data"].startswith(prefixo_busca))
        fix_ano = sum(g["valor"] for g in dados.get("gastos_fixos", []) if g["data"].startswith(prefixo_busca))
        av_ano = sum(g["valor"] for g in dados.get("gastos_avulsos", []) if g["data"].startswith(prefixo_busca))
        gastos_totais = fix_ano + av_ano
        
        # Calcular economias
        economias_mes = 0.0
        for t in dados.get("transferencias", []):
            if t["data"].startswith(prefixo_busca):
                dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
                orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
                if dest_tipo == "Guardado" and orig_tipo == "Normal":
                    economias_mes += t["valor"]
                    
        dados_ano.append({
            "mes_nome": meses_lista[m-1],
            "ganhos": rec_ano,
            "gastos": gastos_totais,
            "guardado": economias_mes,
            "balanco": rec_ano - gastos_totais
        })
        
    total_ganhos_ano = sum(d["ganhos"] for d in dados_ano)
    total_gastos_ano = sum(d["gastos"] for d in dados_ano)
    total_guardado_ano = sum(d["guardado"] for d in dados_ano)
    saldo_anual_acumulado = total_ganhos_ano - total_gastos_ano
    
    doc = docx.Document()
    
    # Título do Relatório
    title = doc.add_paragraph()
    r_title = title.add_run(f"Relatório Financeiro Anual Consolidado - {ano}")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(130, 87, 229)
    
    subtitle = doc.add_paragraph()
    r_sub = subtitle.add_run(f"Finanças Didáticas | Usuário: {usuario_atual}\nGerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M')}")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    
    doc.add_heading("1. Resumo Explicativo do Ano", level=1)
    p1 = doc.add_paragraph()
    p1.add_run(f"O presente relatório consolida o desempenho financeiro do perfil '{usuario_atual}' ao longo do ano de {ano}. A seguir, apresenta-se a síntese geral das movimentações:\n\n")
    p1.add_run("Ganhos Anuais Acumulados: ").bold = True
    p1.add_run(f"O faturamento bruto do ano somou R$ {total_ganhos_ano:,.2f}.\n")
    p1.add_run("Despesas Anuais Acumuladas: ").bold = True
    p1.add_run(f"O montante de saídas reais somou R$ {total_gastos_ano:,.2f}.\n")
    p1.add_run("Balanço Anual: ").bold = True
    p1.add_run(f"O balanço financeiro líquido acumulado foi de R$ {saldo_anual_acumulado:,.2f}. ")
    if total_guardado_ano > 0:
        p1.add_run(f"Deste valor, R$ {total_guardado_ano:,.2f} foram economizados com sucesso nas contas de poupança/investimentos (Guardados).\n")
    else:
        p1.add_run("Não foram registradas transferências de economias direcionadas para a poupança neste ano.\n")
        
    doc.add_heading("2. Detalhamento Mensal (Mês a Mês)", level=1)
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Mês'
    hdr_cells[0].paragraphs[0].runs[0].font.bold = True
    hdr_cells[1].text = 'Ganhos (R$)'
    hdr_cells[1].paragraphs[0].runs[0].font.bold = True
    hdr_cells[2].text = 'Gastos (R$)'
    hdr_cells[2].paragraphs[0].runs[0].font.bold = True
    hdr_cells[3].text = 'Balanço Líquido (R$)'
    hdr_cells[3].paragraphs[0].runs[0].font.bold = True
    
    for d in dados_ano:
        row_cells = table.add_row().cells
        row_cells[0].text = d["mes_nome"]
        row_cells[1].text = f"R$ {d['ganhos']:,.2f}"
        row_cells[2].text = f"R$ {d['gastos']:,.2f}"
        row_cells[3].text = f"R$ {d['balanco']:,.2f}"
        
    # Linha final de total
    total_row = table.add_row().cells
    total_row[0].text = "TOTAL ANUAL"
    total_row[0].paragraphs[0].runs[0].font.bold = True
    total_row[1].text = f"R$ {total_ganhos_ano:,.2f}"
    total_row[1].paragraphs[0].runs[0].font.bold = True
    total_row[2].text = f"R$ {total_gastos_ano:,.2f}"
    total_row[2].paragraphs[0].runs[0].font.bold = True
    total_row[3].text = f"R$ {saldo_anual_acumulado:,.2f}"
    total_row[3].paragraphs[0].runs[0].font.bold = True
    
    for row in table.rows:
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(1.5)
        row.cells[2].width = Inches(1.5)
        row.cells[3].width = Inches(1.5)
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# --- CONTROLE DE MULTIPERFIS E SEGURANÇA ---
def obter_caminho_dados(usuario):
    nome_limpo = usuario.lower().strip().replace(" ", "_")
    return f"dados_financeiros_{nome_limpo}.json"

def carregar_dados_usuario(usuario):
    caminho = obter_caminho_dados(usuario)
    if os.path.exists(caminho):
        try:
            with open(caminho, "r", encoding="utf-8") as f:
                dados = json.load(f)
        except:
            dados = {}
    else:
        dados = {}
        
    # Inicialização padrão de chaves para segurança
    pin_padrao = "123456" if usuario.lower() == "yago" else "654321" if usuario.lower() == "binete" else "123456"
    dados.setdefault("pin", pin_padrao)
    dados.setdefault("receitas", [])
    dados.setdefault("gastos_fixos", [])
    dados.setdefault("gastos_avulsos", [])
    dados.setdefault("gastos_terceiros_cartao", [])
    dados.setdefault("gastos_terceiros_emprestimo", [])
    dados.setdefault("transferencias", [])
    
    # MIGRAR CONTAS ANTIGAS PARA ESTRUTURA DINÂMICA
    if "contas" not in dados:
        saldo_isolada = 0.0
        saldo_caixinha = 0.0
        if "conta_isolada" in dados:
            saldo_isolada = dados["conta_isolada"].get("saldo_principal", 0.0)
            saldo_caixinha = dados["conta_isolada"].get("caixinha_nu", 0.0)
            
        dados["contas"] = [
            {"nome": "Nu", "saldo": 0.0, "tipo": "Normal"},
            {"nome": "PicPay", "saldo": 0.0, "tipo": "Normal"},
            {"nome": "Inter", "saldo": 0.0, "tipo": "Normal"},
            {"nome": "Dinheiro", "saldo": 0.0, "tipo": "Normal"},
            {"nome": "Banco Pan (Guardado)", "saldo": saldo_isolada, "tipo": "Guardado"},
            {"nome": "Caixinha NU (Guardado)", "saldo": saldo_caixinha, "tipo": "Guardado"}
        ]
        
    # MIGRAR CARTOES DE CRÉDITO DINÂMICOS
    if "cartoes" not in dados:
        dados["cartoes"] = [
            {"nome": "Cartão Nu", "fatura": 0.0, "fechamento": 1, "vencimento": 10},
            {"nome": "Cartão PicPay", "fatura": 0.0, "fechamento": 1, "vencimento": 10}
        ]

    # Garante que cartões antigos (criados antes desta versão) tenham fechamento/vencimento
    for c in dados.get("cartoes", []):
        c.setdefault("fechamento", 1)
        c.setdefault("vencimento", 10)

    # Migrar/garantir lista de gastos fixos recorrentes ("para sempre")
    dados.setdefault("gastos_recorrentes", [])
    for rec in dados["gastos_recorrentes"]:
        rec.setdefault("fim", None)
        rec.setdefault("pagamentos", {})
        rec.setdefault("valores_override", {})

    # Migrar/garantir lista de dívidas recorrentes de terceiros no cartão (ex: assinatura dividida)
    dados.setdefault("terceiros_recorrentes", [])
    for rec in dados["terceiros_recorrentes"]:
        rec.setdefault("fim", None)
        rec.setdefault("pagamentos", {})
        rec.setdefault("valores_override", {})

    # Log de recebimentos de terceiros (para reconstrução de extrato mensal)
    dados.setdefault("recebimentos_terceiros", [])

    # Migrar pergunta de segurança (recuperação de senha)
    dados.setdefault("pergunta_seguranca", None)
    dados.setdefault("resposta_hash", None)
        
    # Compatibilidade com faturas de cartão nas despesas antigas
    for r in dados.get("receitas", []):
        r.setdefault("conta", "Nu")
        
    for g in dados.get("gastos_fixos", []):
        g.setdefault("metodo_pagamento", "Saldo")
        g.setdefault("conta", "Nu")
        g.setdefault("cartao_nome", "Cartão Nu")
        
    for a in dados.get("gastos_avulsos", []):
        a.setdefault("metodo_pagamento", "Saldo")
        a.setdefault("conta", "Nu")
        a.setdefault("cartao_nome", "Cartão Nu")
        
    for t in dados.get("gastos_terceiros_cartao", []):
        t.setdefault("cartao_nome", "Cartão Nu")
        
    return dados

def salvar_dados_usuario(usuario, dados):
    caminho = obter_caminho_dados(usuario)
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(dados, f, indent=4, ensure_ascii=False)

# --- AUXILIARES DE MOVIMENTAÇÃO DE SALDO & FATURAS ---
def alterar_saldo(dados, nome_conta, valor, operacao="somar"):
    for conta in dados.setdefault("contas", []):
        if conta["nome"] == nome_conta:
            if operacao == "somar":
                conta["saldo"] += valor
            elif operacao == "subtrair":
                conta["saldo"] -= valor
            return True
    return False

def alterar_fatura(dados, nome_cartao, valor, operacao="somar"):
    for cartao in dados.setdefault("cartoes", []):
        if cartao["nome"] == nome_cartao:
            if operacao == "somar":
                cartao["fatura"] += valor
            elif operacao == "subtrair":
                cartao["fatura"] -= valor
                if cartao["fatura"] < 0:
                    cartao["fatura"] = 0.0
            return True
    return False


def calcular_saldo_conta_no_periodo(dados, nome_conta, periodo_fim):
    """
    Reconstrói o saldo de uma conta até o FIM de um período (AAAA-MM), a partir do
    histórico de lançamentos já registrados (mesma lógica de um extrato bancário:
    saldo atual real, menos tudo que ainda não tinha acontecido até aquele mês).

    OBS: usa a data/'período' nominal de cada lançamento (a data que você escolheu ao
    cadastrar) como a data do movimento no livro-razão.
    """
    conta_obj = next((c for c in dados.get("contas", []) if c["nome"] == nome_conta), None)
    if not conta_obj:
        return 0.0

    def per(data_str):
        return str(data_str)[:7]

    ajuste_total = 0.0        # soma de TODOS os movimentos já conhecidos (qualquer data)
    ajuste_ate_periodo = 0.0  # soma apenas dos movimentos com data <= periodo_fim

    def registrar(valor, data_evento):
        nonlocal ajuste_total, ajuste_ate_periodo
        ajuste_total += valor
        if per(data_evento) <= periodo_fim:
            ajuste_ate_periodo += valor

    for r in dados.get("receitas", []):
        if r.get("conta") == nome_conta:
            registrar(r["valor"], r["data"])

    for g in dados.get("gastos_fixos", []):
        if g.get("metodo_pagamento") == "Saldo" and g.get("conta") == nome_conta and g.get("pago"):
            registrar(-g["valor"], g["data"])

    for rec in dados.get("gastos_recorrentes", []):
        if rec.get("metodo_pagamento") == "Saldo" and rec.get("conta") == nome_conta:
            for p, pago_flag in rec.get("pagamentos", {}).items():
                if pago_flag:
                    v = rec.get("valores_override", {}).get(p, rec["valor"])
                    registrar(-v, f"{p}-01")

    for a in dados.get("gastos_avulsos", []):
        if a.get("metodo_pagamento") == "Saldo" and a.get("conta") == nome_conta:
            registrar(-a["valor"], a["data"])

    for t in dados.get("transferencias", []):
        if t.get("origem") == nome_conta:
            registrar(-t["valor"], t["data"])
        if t.get("destino") == nome_conta:
            registrar(t["valor"], t["data"])

    for e in dados.get("gastos_terceiros_emprestimo", []):
        if e.get("conta_origem") == nome_conta:
            registrar(-e["valor"], e["data"])

    for rt in dados.get("recebimentos_terceiros", []):
        if rt.get("conta_destino") == nome_conta:
            registrar(rt["valor"], rt["data"])

    saldo_base = conta_obj["saldo"] - ajuste_total
    return saldo_base + ajuste_ate_periodo


# --- INICIALIZAÇÃO DO ESTADO DA SESSÃO ---
if "usuario_ativo" not in st.session_state:
    params = st.query_params
    usuario_url = params.get("user", "Yago").title()
    st.session_state.usuario_ativo = usuario_url

if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

usuario_atual = st.session_state.usuario_ativo
dados = carregar_dados_usuario(usuario_atual)

# --- CABEÇALHO DO APP ---
col_logo, col_titulo = st.columns([1, 4])
with col_logo:
    if os.path.exists("app_logo_financeiro.png"):
        st.image("app_logo_financeiro.png", width=70)
    else:
        st.markdown("<h1 style='font-size: 50px; margin: 0;'>💰</h1>", unsafe_allow_html=True)

with col_titulo:
    st.title("Finanças Didáticas")
    st.caption("Controle financeiro premium e descomplicado")

# --- TELA DE LOGIN / SEGURANÇA ---
if not st.session_state.autenticado:
    st.markdown(f"### 🔒 Acesso Restrito - Perfil: **{usuario_atual}**")
    st.markdown("Por segurança, digite sua senha de 6 dígitos para acessar seus dados.")
    
    tentativas_key = f"tentativas_login_{usuario_atual}"
    if tentativas_key not in st.session_state:
        st.session_state[tentativas_key] = 0

    bloqueado = st.session_state[tentativas_key] >= 5

    if bloqueado:
        st.error("🚫 Muitas tentativas incorretas. Recarregue a página para tentar novamente.")

    pin_digitado = st.text_input("Digite o PIN de 6 dígitos:", type="password", max_chars=6, disabled=bloqueado)
    
    col_btn_login, col_btn_perfil = st.columns(2)
    with col_btn_login:
        if st.button("Entrar 🔓", disabled=bloqueado):
            if pin_digitado == dados.get("pin", "123456"):
                st.session_state.autenticado = True
                st.session_state[tentativas_key] = 0
                st.success("Acesso liberado!")
                st.rerun()
            else:
                st.session_state[tentativas_key] += 1
                restantes = 5 - st.session_state[tentativas_key]
                if restantes > 0:
                    st.error(f"PIN incorreto! Tentativas restantes: {restantes}.")
                else:
                    st.rerun()
                
    with col_btn_perfil:
        novo_usuario = st.selectbox("Trocar de Perfil:", ["Yago", "Binete", "Criar Novo Perfil"])
        if novo_usuario == "Criar Novo Perfil":
            st.markdown("---")
            st.markdown("**Criar Novo Perfil Independente**")
            novo_nome = st.text_input("Nome do Novo Usuário:")
            novo_pin = st.text_input("Escolha um PIN de 6 dígitos (somente números):", type="password", max_chars=6)

            pergunta_escolhida = st.selectbox("Escolha uma pergunta de segurança (para recuperar o PIN se esquecer):", PERGUNTAS_SEGURANCA_PADRAO)
            pergunta_final = pergunta_escolhida
            if pergunta_escolhida == PERGUNTAS_SEGURANCA_PADRAO[-1]:
                pergunta_final = st.text_input("Digite sua pergunta personalizada:")
            resposta_seguranca = st.text_input("Resposta da pergunta de segurança:")

            if st.button("Criar Perfil e Acessar 🚀"):
                if novo_nome and len(novo_pin) == 6 and novo_pin.isdigit() and pergunta_final and resposta_seguranca:
                    dados_novos = carregar_dados_usuario(novo_nome)
                    dados_novos["pin"] = novo_pin
                    dados_novos["pergunta_seguranca"] = pergunta_final
                    dados_novos["resposta_hash"] = hash_resposta(resposta_seguranca)
                    salvar_dados_usuario(novo_nome, dados_novos)
                    st.session_state.usuario_ativo = novo_nome.title()
                    st.session_state.autenticado = True
                    st.query_params["user"] = novo_nome.lower()
                    st.success(f"Perfil de {novo_nome} criado e logado!")
                    st.rerun()
                else:
                    st.error("Preencha nome, PIN de 6 dígitos, a pergunta e a resposta de segurança.")
        elif novo_usuario != usuario_atual:
            st.session_state.usuario_ativo = novo_usuario
            st.query_params["user"] = novo_usuario.lower()
            st.rerun()

    # --- ESQUECI MEU PIN (RECUPERAÇÃO POR PERGUNTA DE SEGURANÇA) ---
    st.markdown("---")
    with st.expander("🔑 Esqueci meu PIN"):
        if not dados.get("pergunta_seguranca") or not dados.get("resposta_hash"):
            st.info(
                "Este perfil ainda não tem uma pergunta de segurança configurada. "
                "Faça login normalmente e configure uma na aba ⚙️ Senha para poder recuperar o acesso no futuro."
            )
        else:
            st.write(f"**Pergunta de segurança:** {dados['pergunta_seguranca']}")
            resposta_tentativa = st.text_input("Sua resposta:", key="resposta_recuperacao")
            novo_pin_rec = st.text_input("Novo PIN (6 dígitos):", type="password", max_chars=6, key="novo_pin_rec")
            confirmar_pin_rec = st.text_input("Confirme o novo PIN:", type="password", max_chars=6, key="conf_pin_rec")
            if st.button("Redefinir PIN"):
                if hash_resposta(resposta_tentativa) != dados["resposta_hash"]:
                    st.error("Resposta incorreta.")
                elif not (novo_pin_rec.isdigit() and len(novo_pin_rec) == 6):
                    st.error("O novo PIN deve ter exatamente 6 números.")
                elif novo_pin_rec != confirmar_pin_rec:
                    st.error("Os dois PINs digitados não coincidem.")
                else:
                    dados["pin"] = novo_pin_rec
                    salvar_dados_usuario(usuario_atual, dados)
                    st.session_state[tentativas_key] = 0
                    st.success("PIN redefinido com sucesso! Você já pode entrar com o novo PIN.")
                    st.rerun()

    st.stop()

# --- SELETOR DE MÊS E ANO (PLANEJAMENTO SEM LIMITES) ---
st.markdown("---")
col_perf, col_m, col_a = st.columns([2, 2, 2])

with col_perf:
    st.markdown(f"👤 **{usuario_atual}**")
    if st.button("Sair / Bloquear 🔒"):
        st.session_state.autenticado = False
        st.rerun()

with col_m:
    meses_lista = [
        "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
        "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"
    ]
    mes_atual_nome = meses_lista[datetime.now().month - 1]
    mes_selecionado = st.selectbox("Selecione o Mês:", meses_lista, index=datetime.now().month - 1)
    mes_num = meses_lista.index(mes_selecionado) + 1

with col_a:
    anos_lista = list(range(2025, 2036))
    ano_selecionado = st.selectbox("Selecione o Ano:", anos_lista, index=anos_lista.index(datetime.now().year))

periodo_ativo = f"{ano_selecionado}-{mes_num:02d}"

# --- NAVEGAÇÃO POR ABAS (Mobile Friendly) ---
abas = st.tabs(["📊 Resumo", "💵 Ganhos", "🏠 Fixas", "🛍️ Avulsos", "👥 Terceiros", "💳 Saldos", "🔒 Guardado", "⚙️ Senha"])

# ----------------- ABA 1: RESUMO GERAL & RELATÓRIOS -----------------
with abas[0]:
    st.markdown(f"<h3 class='titulo-secao'>📊 Painel Financeiro ({mes_selecionado}/{ano_selecionado})</h3>", unsafe_allow_html=True)
    
    saldo_normal_contas = sum(
        calcular_saldo_conta_no_periodo(dados, c["nome"], periodo_ativo)
        for c in dados.get("contas", []) if c.get("tipo", "Normal") == "Normal"
    )
    saldo_guardado_contas = sum(
        calcular_saldo_conta_no_periodo(dados, c["nome"], periodo_ativo)
        for c in dados.get("contas", []) if c.get("tipo", "Normal") == "Guardado"
    )
    faturas_abertas = sum(c["fatura"] for c in dados.get("cartoes", []))
    patrimonio_liquido = saldo_normal_contas + saldo_guardado_contas - faturas_abertas

    col_pl1, col_pl2 = st.columns(2)
    col_pl1.metric(f"🏦 Contas Normais ({mes_selecionado}/{ano_selecionado})", f"R$ {saldo_normal_contas:,.2f}")
    col_pl2.metric(f"🔒 Guardado/Poupança ({mes_selecionado}/{ano_selecionado})", f"R$ {saldo_guardado_contas:,.2f}")

    col_pl3, col_pl4 = st.columns(2)
    col_pl3.metric("Faturas em Aberto (hoje)", f"R$ {faturas_abertas:,.2f}")
    col_pl4.metric("Patrimônio Líquido", f"R$ {patrimonio_liquido:,.2f}")
    st.caption("Os saldos de contas refletem o mês/ano selecionado no topo. As faturas mostram o valor real de hoje.")
    st.markdown("---")


    tipo_relatorio = st.radio("Escolha o tipo de Relatório:", ["Relatório Mensal", "Relatório Anual"], horizontal=True)
    
    if tipo_relatorio == "Relatório Mensal":
        # Filtrar dados do mês selecionado
        receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo_ativo)]
        gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if pertence_periodo_cartao(g, periodo_ativo)]
        gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if pertence_periodo_cartao(g, periodo_ativo)]
        
        total_receitas = sum(r["valor"] for r in receitas_mes)
        total_fixos = sum(g["valor"] for g in gastos_fixos_mes)
        total_avulsos = sum(g["valor"] for g in gastos_avulsos_mes)
        total_gastos = total_fixos + total_avulsos
        saldo_livre = total_receitas - total_gastos
        
        # Calcular total guardado nas contas do tipo "Guardado"
        total_guardado = sum(c["saldo"] for c in dados["contas"] if c["tipo"] == "Guardado")
        
        # Terceiros pendentes
        unpaid_cartao = [t for t in dados.get("gastos_terceiros_cartao", []) if not t["pago"]]
        unpaid_emprestimo = [e for e in dados.get("gastos_terceiros_emprestimo", []) if not e["pago"]]
        
        total_cartao_terceiros = sum(t["valor"] for t in unpaid_cartao if pertence_periodo_cartao(t, periodo_ativo))
        total_emprestimos_pendentes = sum(e["valor"] for e in unpaid_emprestimo)
        
        # Métricas na tela
        col_m1, col_m2 = st.columns(2)
        with col_m1:
            st.metric(label="Minhas Receitas (+)", value=f"R$ {total_receitas:,.2f}")
            st.metric(label="Meus Gastos (-)", value=f"R$ {total_gastos:,.2f}")
        with col_m2:
            st.metric(label="Saldo do Mês (Livre)", value=f"R$ {saldo_livre:,.2f}", delta=f"R$ {saldo_livre:,.2f}" if saldo_livre >= 0 else f"R$ {saldo_livre:,.2f}")
            st.metric(label="Total Guardado", value=f"R$ {total_guardado:,.2f}")
            
        # Alertas de Terceiros
        st.markdown("#### 👥 Dinheiro de Terceiros a Receber:")
        c_t1, c_t2 = st.columns(2)
        with c_t1:
            if total_cartao_terceiros > 0:
                st.warning(f"💳 **Cartão de Crédito:**\n\nR$ {total_cartao_terceiros:,.2f} pendentes neste mês.")
            else:
                st.success("💳 **Cartão de Crédito:**\n\nNenhum reembolso pendente.")
        with c_t2:
            if total_emprestimos_pendentes > 0:
                st.warning(f"💸 **Empréstimos:**\n\nR$ {total_emprestimos_pendentes:,.2f} a receber.")
            else:
                st.success("💸 **Empréstimos:**\n\nTodos pagos!")
                
        # Gráficos do Mês (Container para Empilhar no Celular e Colunas no PC)
        cont_graf = st.container()
        with cont_graf:
            # Gráfico 1: Divisão das Despesas (Pizza)
            if total_gastos > 0:
                st.markdown("#### 📈 Divisão das Despesas (Fixos vs. Avulsos)")
                df_pizza = pd.DataFrame([
                    {"Categoria": "Gastos Fixos", "Valor": total_fixos},
                    {"Categoria": "Gastos Avulsos", "Valor": total_avulsos}
                ])
                fig = px.pie(df_pizza, values="Valor", names="Categoria", hole=0.4,
                             color_discrete_sequence=["#8257E5", "#FF4757"])
                fig.update_layout(
                    paper_bgcolor='rgba(0,0,0,0)',
                    plot_bgcolor='rgba(0,0,0,0)',
                    font_color="#FFFFFF",
                    margin=dict(t=10, b=10, l=10, r=10),
                    height=220
                )
                st.plotly_chart(fig, use_container_width=True)
                
            # Gráfico 2: Alocação Geral Macro (Ganhos vs. Despesas vs. Guardados do Mês)
            economias_mes = 0.0
            for t in dados.get("transferencias", []):
                if t["data"].startswith(periodo_ativo):
                    dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
                    orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
                    if dest_tipo == "Guardado" and orig_tipo == "Normal":
                        economias_mes += t["valor"]
                        
            st.markdown("#### 📊 Alocação Macro Mensal (Fluxos)")
            df_macro = pd.DataFrame([
                {"Fluxo": "Ganhos", "Valor": total_receitas, "Cor": "#26DE81"},
                {"Fluxo": "Despesas", "Valor": total_gastos, "Cor": "#FF4757"},
                {"Fluxo": "Guardado", "Valor": economias_mes, "Cor": "#8257E5"}
            ])
            fig_macro = px.bar(df_macro, x="Valor", y="Fluxo", color="Fluxo", orientation="h",
                               color_discrete_map={"Ganhos": "#26DE81", "Despesas": "#FF4757", "Guardado": "#8257E5"})
            fig_macro.update_layout(
                paper_bgcolor='rgba(0,0,0,0)',
                plot_bgcolor='rgba(0,0,0,0)',
                font_color="#FFFFFF",
                margin=dict(t=10, b=10, l=10, r=10),
                height=200,
                showlegend=False
            )
            st.plotly_chart(fig_macro, use_container_width=True)
            
        # --- EXPORTAR RELATÓRIO MENSAL (BOTÕES) ---
        st.markdown("---")
        st.markdown("### 📥 Exportar Relatório Oficial (Mensal)")
        st.caption("Gere um relatório explicativo completo, pronto para baixar em PDF ou editar em Word.")
        
        col_pdf, col_docx = st.columns(2)
        with col_pdf:
            pdf_bytes = gerar_relatorio_mensal_pdf(dados, periodo_ativo, mes_selecionado, ano_selecionado, usuario_atual)
            st.download_button(
                label="📄 Baixar PDF",
                data=pdf_bytes,
                file_name=f"relatorio_financeiro_{usuario_atual.lower()}_{periodo_ativo}.pdf",
                mime="application/pdf"
            )
        with col_docx:
            docx_bytes = gerar_relatorio_mensal_docx(dados, periodo_ativo, mes_selecionado, ano_selecionado, usuario_atual)
            st.download_button(
                label="📝 Baixar Word (.docx)",
                data=docx_bytes,
                file_name=f"relatorio_financeiro_{usuario_atual.lower()}_{periodo_ativo}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
    else:
        # --- RELATÓRIO ANUAL CONSOLIDADO ---
        st.markdown(f"### 📅 Relatório Anual - Ano: **{ano_selecionado}**")
        
        # Consolidar dados de todos os 12 meses do ano selecionado de forma correta
        dados_ano = []
        for m in range(1, 13):
            prefixo_busca = f"{ano_selecionado}-{m:02d}"
            rec_ano = sum(r["valor"] for r in dados.get("receitas", []) if r["data"].startswith(prefixo_busca))
            fix_ano = sum(g["valor"] for g in dados.get("gastos_fixos", []) if g["data"].startswith(prefixo_busca))
            av_ano = sum(g["valor"] for g in dados.get("gastos_avulsos", []) if g["data"].startswith(prefixo_busca))
            gastos_totais = fix_ano + av_ano
            dados_ano.append({
                "Mês": meses_lista[m-1][:3],
                "Ganhos": rec_ano,
                "Gastos": gastos_totais
            })
            
        df_ano = pd.DataFrame(dados_ano)
        
        # Totais Anuais
        total_ganhos_ano = df_ano["Ganhos"].sum()
        total_gastos_ano = df_ano["Gastos"].sum()
        saldo_anual_acumulado = total_ganhos_ano - total_gastos_ano
        
        col_an1, col_an2, col_an3 = st.columns(3)
        col_an1.metric("Ganhos no Ano", f"R$ {total_ganhos_ano:,.2f}")
        col_an2.metric("Gastos no Ano", f"R$ {total_gastos_ano:,.2f}")
        col_an3.metric("Balanço Acumulado", f"R$ {saldo_anual_acumulado:,.2f}", 
                        delta=f"R$ {saldo_anual_acumulado:,.2f}" if saldo_anual_acumulado >= 0 else f"R$ {saldo_anual_acumulado:,.2f}")
        
        # Gráfico de Linha/Tendência Premium
        st.markdown("#### 📈 Tendência Anual (Ganhos vs. Gastos)")
        df_melted = df_ano.melt(id_vars="Mês", value_vars=["Ganhos", "Gastos"], var_name="Categoria", value_name="Valor (R$)")
        fig_trend = px.line(df_melted, x="Mês", y="Valor (R$)", color="Categoria", markers=True,
                            color_discrete_map={"Ganhos": "#26DE81", "Gastos": "#FF4757"})
        fig_trend.update_layout(
            paper_bgcolor='rgba(0,0,0,0)',
            plot_bgcolor='rgba(0,0,0,0)',
            font_color="#FFFFFF",
            margin=dict(t=10, b=10, l=10, r=10),
            height=280
        )
        st.plotly_chart(fig_trend, use_container_width=True)
        
        # --- EXPORTAR RELATÓRIO ANUAL (BOTÕES) ---
        st.markdown("---")
        st.markdown("### 📥 Exportar Relatório Oficial (Anual)")
        st.caption("Gere o consolidado completo de todos os meses do ano em PDF ou Word.")
        
        col_pdf_an, col_docx_an = st.columns(2)
        with col_pdf_an:
            pdf_an_bytes = gerar_relatorio_anual_pdf(dados, ano_selecionado, meses_lista, usuario_atual)
            st.download_button(
                label="📄 Baixar PDF Anual",
                data=pdf_an_bytes,
                file_name=f"relatorio_anual_{usuario_atual.lower()}_{ano_selecionado}.pdf",
                mime="application/pdf"
            )
        with col_docx_an:
            docx_an_bytes = gerar_relatorio_anual_docx(dados, ano_selecionado, meses_lista, usuario_atual)
            st.download_button(
                label="📝 Baixar Word Anual (.docx)",
                data=docx_an_bytes,
                file_name=f"relatorio_anual_{usuario_atual.lower()}_{ano_selecionado}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

# ----------------- ABA 2: GANHOS (RECEITAS) -----------------
with abas[1]:
    st.markdown("<h3 class='titulo-secao'>💵 Registro de Ganhos</h3>", unsafe_allow_html=True)
    
    with st.form("form_receita", clear_on_submit=True):
        st.markdown("**Adicionar Nova Receita**")
        desc = st.text_input("Descrição (Ex: Salário, PIX, Freelance)")
        val = st.number_input("Valor (R$)", min_value=0.0, step=50.0, format="%.2f")
        
        # Selecionar em qual conta depositar
        contas_nomes = [c["nome"] for c in dados["contas"]]
        conta_dest = st.selectbox("Depositar na Conta:", contas_nomes)
        
        data_padrao = datetime(ano_selecionado, mes_num, min(datetime.now().day, 28))
        data_rec = st.date_input("Data do Recebimento", data_padrao, format="DD/MM/YYYY")
        
        enviar = st.form_submit_button("Salvar Receita")
        if enviar and desc and val > 0:
            nova_rec = {"data": str(data_rec), "descricao": desc, "valor": val, "conta": conta_dest}
            dados.setdefault("receitas", []).append(nova_rec)
            
            # Atualiza saldo automaticamente
            alterar_saldo(dados, conta_dest, val, "somar")
            
            salvar_dados_usuario(usuario_atual, dados)
            st.success(f"Receita adicionada e valor depositado na conta {conta_dest}!")
            st.rerun()

    # Mostrar Receitas do Mês Selecionado
    receitas_mes = [r for r in dados.get("receitas", []) if r["data"].startswith(periodo_ativo)]
    if receitas_mes:
        st.markdown("#### Suas Receitas Registradas neste Mês")

        filtro_rec = st.text_input("🔎 Buscar por descrição:", key="filtro_receitas")

        for idx, item in enumerate(dados["receitas"]):
            if not item["data"].startswith(periodo_ativo):
                continue
            if filtro_rec and filtro_rec.lower() not in item["descricao"].lower():
                continue

            col_d, col_v, col_del = st.columns([3, 2, 1])
            col_d.markdown(f"**{item['descricao']}**\n\n*{formatar_data_br(item['data'])} — {item.get('conta', 'Não especificada')}*")
            col_v.markdown(f"R$ {item['valor']:,.2f}")

            confirm_key = f"confirma_del_rec_{idx}"
            if col_del.button("🗑️", key=f"del_rec_{idx}", help="Excluir este lançamento"):
                st.session_state[confirm_key] = True

            if st.session_state.get(confirm_key):
                st.warning(f"Excluir '{item['descricao']}' (R$ {item['valor']:,.2f})? Isso também estorna o valor da conta '{item.get('conta', '')}'.")
                c_sim, c_nao = st.columns(2)
                if c_sim.button("Sim, excluir", key=f"conf_sim_rec_{idx}"):
                    alterar_saldo(dados, item.get("conta", ""), item["valor"], "subtrair")
                    dados["receitas"].pop(idx)
                    salvar_dados_usuario(usuario_atual, dados)
                    del st.session_state[confirm_key]
                    st.rerun()
                if c_nao.button("Cancelar", key=f"conf_nao_rec_{idx}"):
                    del st.session_state[confirm_key]
                    st.rerun()
            st.markdown("---")

        with st.expander("⚠️ Apagar todo o histórico de receitas deste mês"):
            st.caption("Atenção: isso remove todos os lançamentos do mês, sem estornar os saldos.")
            confirma_limpar_rec = st.checkbox("Confirmo que quero apagar tudo", key="chk_limpar_rec")
            if st.button("Limpar Histórico de Receitas (Mês)", key="limpar_rec", disabled=not confirma_limpar_rec):
                dados["receitas"] = [r for r in dados.get("receitas", []) if not r["data"].startswith(periodo_ativo)]
                salvar_dados_usuario(usuario_atual, dados)
                st.rerun()

# ----------------- ABA 3: MEUS GASTOS FIXOS -----------------
with abas[2]:
    st.markdown("<h3 class='titulo-secao'>🏠 Meus Gastos Fixos & Parcelas</h3>", unsafe_allow_html=True)
    st.caption("Gastos repetitivos ou parcelados de longo prazo.")

    if "fixo_form_key" not in st.session_state:
        st.session_state.fixo_form_key = 0
    fk = st.session_state.fixo_form_key

    st.markdown("**Adicionar Gasto Fixo**")
    desc = st.text_input("Descrição do Gasto (Ex: Aluguel, Parcela de Notebook, Academia)", key=f"desc_fixo_{fk}")
    val = st.number_input("Valor Mensal (R$)", min_value=0.0, step=10.0, format="%.2f", key=f"val_fixo_{fk}")

    # INTEGRAÇÃO CARTÃO DE CRÉDITO DINÂMICO (fora do form para reagir na hora)
    metodo_p = st.selectbox("Forma de Pagamento:", ["Saldo em Conta", "Cartão de Crédito"], key=f"metodo_fixo_{fk}")
    if metodo_p == "Saldo em Conta":
        conta_pagamento = st.selectbox("Pagar com a Conta:", [c["nome"] for c in dados["contas"]], key=f"conta_fixo_{fk}")
        cartao_pagamento = "Não se aplica"
    else:
        conta_pagamento = "Não se aplica"
        cartao_pagamento = st.selectbox("Pagar com o Cartão:", [c["nome"] for c in dados["cartoes"]], key=f"cartao_fixo_{fk}")

    tipo_despesa_fixa = st.radio(
        "Tipo de Despesa:",
        ["Única (só este mês)", "Parcelada (número fixo de parcelas)", "Recorrente (todo mês, até eu encerrar)"],
        key=f"tipo_fixo_{fk}"
    )
    st.caption("Use 'Recorrente' para gastos como academia, assinaturas ou aluguel: ela continua aparecendo todo mês até você encerrar.")

    total_parc = 1
    if tipo_despesa_fixa == "Parcelada (número fixo de parcelas)":
        total_parc = st.number_input("Número total de parcelas:", min_value=2, max_value=48, value=2, step=1, key=f"parc_fixo_{fk}")

    pago = st.checkbox("Marcar já como pago neste mês?", key=f"pago_fixo_{fk}")

    if st.button("Salvar Despesa", key=f"salvar_fixo_{fk}"):
        if desc and val > 0:
            data_inicial = datetime(ano_selecionado, mes_num, 1)
            metodo_salvar = "Saldo" if metodo_p == "Saldo em Conta" else "Cartao"

            if tipo_despesa_fixa == "Recorrente (todo mês, até eu encerrar)":
                novo_recorrente = {
                    "id": f"rec_{datetime.now().strftime('%Y%m%d%H%M%S%f')}",
                    "descricao": desc,
                    "valor": val,
                    "metodo_pagamento": metodo_salvar,
                    "conta": conta_pagamento,
                    "cartao_nome": cartao_pagamento,
                    "inicio": periodo_ativo,
                    "fim": None,
                    "pagamentos": {periodo_ativo: pago},
                    "valores_override": {}
                }
                dados.setdefault("gastos_recorrentes", []).append(novo_recorrente)
                if pago:
                    if metodo_salvar == "Saldo":
                        alterar_saldo(dados, conta_pagamento, val, "subtrair")
                    else:
                        alterar_fatura(dados, cartao_pagamento, val, "somar")

            elif tipo_despesa_fixa == "Parcelada (número fixo de parcelas)":
                lista_periodos = get_proximos_meses(data_inicial, total_parc)
                for i, periodo in enumerate(lista_periodos):
                    status_pago_parc = pago if i == 0 else False
                    novo_gasto = {
                        "data": f"{periodo}-01",
                        "descricao": f"{desc} (Parc. {i+1}/{total_parc})",
                        "valor": val,
                        "pago": status_pago_parc,
                        "metodo_pagamento": metodo_salvar,
                        "conta": conta_pagamento,
                        "cartao_nome": cartao_pagamento,
                        "periodo_fatura": periodo
                    }
                    dados.setdefault("gastos_fixos", []).append(novo_gasto)
                    if status_pago_parc:
                        if metodo_salvar == "Saldo":
                            alterar_saldo(dados, conta_pagamento, val, "subtrair")
                        else:
                            alterar_fatura(dados, cartao_pagamento, val, "somar")
            else:
                periodo_fat_fixo = periodo_ativo
                if metodo_salvar == "Cartao":
                    cartao_obj_fixo = obter_dados_cartao(dados, cartao_pagamento)
                    if cartao_obj_fixo:
                        periodo_fat_fixo = calcular_periodo_fatura(f"{periodo_ativo}-01", cartao_obj_fixo.get("fechamento", 1))
                novo_gasto = {
                    "data": f"{periodo_ativo}-01",
                    "descricao": desc,
                    "valor": val,
                    "pago": pago,
                    "metodo_pagamento": metodo_salvar,
                    "conta": conta_pagamento,
                    "cartao_nome": cartao_pagamento,
                    "periodo_fatura": periodo_fat_fixo
                }
                dados.setdefault("gastos_fixos", []).append(novo_gasto)
                if pago:
                    if metodo_salvar == "Saldo":
                        alterar_saldo(dados, conta_pagamento, val, "subtrair")
                    else:
                        alterar_fatura(dados, cartao_pagamento, val, "somar")
                
            salvar_dados_usuario(usuario_atual, dados)
            st.session_state.fixo_form_key += 1
            st.success("Despesa cadastrada com sucesso!")
            st.rerun()
        else:
            st.error("Preencha a descrição e um valor maior que zero.")

    # Mostrar Gastos Recorrentes ativos neste mês
    recorrentes_ativos = [
        r for r in dados.get("gastos_recorrentes", [])
        if r["inicio"] <= periodo_ativo and (r["fim"] is None or periodo_ativo < r["fim"])
    ]
    if recorrentes_ativos:
        st.markdown("#### 🔁 Despesas Recorrentes (Todo Mês)")
        for rec in recorrentes_ativos:
            valor_mes_rec = rec["valores_override"].get(periodo_ativo, rec["valor"])
            pago_mes_rec = rec["pagamentos"].get(periodo_ativo, False)
            rec_metodo = rec.get("metodo_pagamento", "Saldo")
            rec_local = rec.get("conta", "Nu") if rec_metodo == "Saldo" else rec.get("cartao_nome", "Cartão Nu")

            col_d, col_v, col_p = st.columns([2, 1, 1])
            col_d.markdown(f"**{rec['descricao']}** 🔁\n\n*({rec_local})*")
            col_v.markdown(f"R$ {valor_mes_rec:,.2f}")
            novo_pago_rec = col_p.checkbox("Pago", value=pago_mes_rec, key=f"rec_pago_{rec['id']}_{periodo_ativo}")

            if novo_pago_rec != pago_mes_rec:
                if novo_pago_rec:
                    if rec_metodo == "Saldo":
                        alterar_saldo(dados, rec_local, valor_mes_rec, "subtrair")
                    else:
                        alterar_fatura(dados, rec_local, valor_mes_rec, "somar")
                else:
                    if rec_metodo == "Saldo":
                        alterar_saldo(dados, rec_local, valor_mes_rec, "somar")
                    else:
                        alterar_fatura(dados, rec_local, valor_mes_rec, "subtrair")
                rec["pagamentos"][periodo_ativo] = novo_pago_rec
                salvar_dados_usuario(usuario_atual, dados)
                st.rerun()

            with st.expander(f"⚙️ Gerenciar '{rec['descricao']}'"):
                novo_valor_rec = st.number_input(
                    f"Valor de {mes_selecionado}/{ano_selecionado}:", min_value=0.0,
                    value=float(valor_mes_rec), step=5.0, format="%.2f", key=f"editval_rec_{rec['id']}"
                )
                if st.button("Salvar novo valor deste mês", key=f"btn_editval_rec_{rec['id']}"):
                    diff_rec = novo_valor_rec - valor_mes_rec
                    if pago_mes_rec and diff_rec != 0:
                        if rec_metodo == "Saldo":
                            alterar_saldo(dados, rec_local, diff_rec, "subtrair")
                        else:
                            alterar_fatura(dados, rec_local, diff_rec, "somar")
                    rec["valores_override"][periodo_ativo] = novo_valor_rec
                    salvar_dados_usuario(usuario_atual, dados)
                    st.rerun()

                st.markdown("---")
                st.caption("Encerrar esta recorrência: ela deixa de aparecer a partir do mês escolhido (meses anteriores continuam no histórico).")
                if st.button(f"🛑 Encerrar a partir de {mes_selecionado}/{ano_selecionado}", key=f"encerrar_rec_{rec['id']}"):
                    rec["fim"] = periodo_ativo
                    salvar_dados_usuario(usuario_atual, dados)
                    st.success("Recorrência encerrada a partir deste mês.")
                    st.rerun()

                st.markdown("---")
                confirma_del_rec = st.checkbox("Confirmo que quero excluir esta recorrência por completo", key=f"chk_del_rec_{rec['id']}")
                if st.button("🗑️ Excluir recorrência por completo", key=f"del_rec_{rec['id']}", disabled=not confirma_del_rec):
                    dados["gastos_recorrentes"] = [r for r in dados["gastos_recorrentes"] if r["id"] != rec["id"]]
                    salvar_dados_usuario(usuario_atual, dados)
                    st.rerun()
            st.markdown("---")

    # Mostrar recorrências encerradas que poderiam ser reativadas (apenas informativo, dentro de um expander)
    recorrentes_encerradas = [r for r in dados.get("gastos_recorrentes", []) if r["fim"] is not None and r["fim"] <= periodo_ativo]
    if recorrentes_encerradas:
        with st.expander("↩️ Recorrências encerradas (reativar se necessário)"):
            for rec in recorrentes_encerradas:
                col_re1, col_re2 = st.columns([3, 1])
                col_re1.markdown(f"**{rec['descricao']}** — encerrada a partir de {rec['fim']}")
                if col_re2.button("Reativar", key=f"reativar_{rec['id']}"):
                    rec["fim"] = None
                    salvar_dados_usuario(usuario_atual, dados)
                    st.rerun()

    # Mostrar Gastos Fixos do Mês (únicos e parcelados)
    gastos_fixos_mes = [g for g in dados.get("gastos_fixos", []) if pertence_periodo_cartao(g, periodo_ativo)]
    if gastos_fixos_mes:
        st.markdown("#### Controle de Pagamentos de Despesas Fixas / Parceladas")
        for idx, item in enumerate(dados["gastos_fixos"]):
            if pertence_periodo_cartao(item, periodo_ativo):
                col_d, col_v, col_p = st.columns([2, 1, 1])
                item_metodo = item.get("metodo_pagamento", "Saldo")
                item_local = item.get("conta", "Nu") if item_metodo == "Saldo" else item.get("cartao_nome", "Cartão Nu")
                col_d.markdown(f"**{item['descricao']}**\n\n*({item_local})*")
                col_v.markdown(f"R$ {item['valor']:,.2f}")
                
                status_pago = col_p.checkbox("Pago", value=item["pago"], key=f"fixo_{idx}_{periodo_ativo}")
                if status_pago != item["pago"]:
                    # Atualiza saldo ou fatura de cartão
                    if status_pago: # Marcou pago agora
                        if item_metodo == "Saldo":
                            alterar_saldo(dados, item_local, item["valor"], "subtrair")
                        else:
                            alterar_fatura(dados, item_local, item["valor"], "somar")
                    else: # Desmarcou pagamento, estorna
                        if item_metodo == "Saldo":
                            alterar_saldo(dados, item_local, item["valor"], "somar")
                        else:
                            alterar_fatura(dados, item_local, item["valor"], "subtrair")
                        
                    dados["gastos_fixos"][idx]["pago"] = status_pago
                    salvar_dados_usuario(usuario_atual, dados)
                    st.rerun()

                with st.expander(f"⚙️ Editar / Excluir '{item['descricao']}'"):
                    novo_valor_fixo = st.number_input(
                        "Valor desta parcela:", min_value=0.0, value=float(item["valor"]),
                        step=5.0, format="%.2f", key=f"editval_fixo_{idx}"
                    )
                    if st.button("Salvar novo valor", key=f"btn_editval_fixo_{idx}"):
                        diff_fixo = novo_valor_fixo - item["valor"]
                        if item["pago"] and diff_fixo != 0:
                            if item_metodo == "Saldo":
                                alterar_saldo(dados, item_local, diff_fixo, "subtrair")
                            else:
                                alterar_fatura(dados, item_local, diff_fixo, "somar")
                        dados["gastos_fixos"][idx]["valor"] = novo_valor_fixo
                        salvar_dados_usuario(usuario_atual, dados)
                        st.rerun()

                    st.markdown("---")
                    confirma_del_fixo = st.checkbox("Confirmo que quero excluir este item", key=f"chk_del_fixo_{idx}")
                    if st.button("🗑️ Excluir este item", key=f"del_fixo_{idx}", disabled=not confirma_del_fixo):
                        if item["pago"]:
                            if item_metodo == "Saldo":
                                alterar_saldo(dados, item_local, item["valor"], "somar")
                            else:
                                alterar_fatura(dados, item_local, item["valor"], "subtrair")
                        dados["gastos_fixos"].pop(idx)
                        salvar_dados_usuario(usuario_atual, dados)
                        st.rerun()
                st.markdown("---")
                    
        with st.expander("⚠️ Apagar todos os gastos fixos deste mês"):
            st.caption("Atenção: isso remove todos os lançamentos do mês, sem estornar os saldos/faturas.")
            confirma_limpar_fixos = st.checkbox("Confirmo que quero apagar tudo", key="chk_limpar_fixos")
            if st.button("Limpar Todos os Gastos Fixos (Deste Mês)", key="limpar_fixos", disabled=not confirma_limpar_fixos):
                dados["gastos_fixos"] = [g for g in dados.get("gastos_fixos", []) if not pertence_periodo_cartao(g, periodo_ativo)]
                salvar_dados_usuario(usuario_atual, dados)
                st.rerun()

# ----------------- ABA 4: GASTOS AVULSOS -----------------
with abas[3]:
    st.markdown("<h3 class='titulo-secao'>🛍️ Meus Gastos Avulsos</h3>", unsafe_allow_html=True)
    st.caption("Gastos diários variáveis (Ifood, Uber, Compras rápidas)")

    if "avulso_form_key" not in st.session_state:
        st.session_state.avulso_form_key = 0
    fka = st.session_state.avulso_form_key

    st.markdown("**Adicionar Gasto Avulso**")
    desc = st.text_input("O que você comprou?", key=f"desc_av_{fka}")
    val = st.number_input("Valor Pago (R$)", min_value=0.0, step=10.0, format="%.2f", key=f"val_av_{fka}")
    
    # INTEGRAÇÃO CARTÃO DE CRÉDITO DINÂMICO
    metodo_p_av = st.selectbox("Forma de Pagamento:", ["Saldo em Conta", "Cartão de Crédito"], key=f"metodo_av_{fka}")
    if metodo_p_av == "Saldo em Conta":
        conta_pagamento_av = st.selectbox("Pagar com a Conta:", [c["nome"] for c in dados["contas"]], key=f"conta_av_{fka}")
        cartao_pagamento_av = "Não se aplica"
    else:
        conta_pagamento_av = "Não se aplica"
        cartao_pagamento_av = st.selectbox("Pagar com o Cartão:", [c["nome"] for c in dados["cartoes"]], key=f"cartao_av_{fka}")
    
    data_padrao = datetime(ano_selecionado, mes_num, min(datetime.now().day, 28))
    data_gasto = st.date_input("Data do Gasto", data_padrao, format="DD/MM/YYYY", key=f"data_av_{fka}")
    
    if st.button("Salvar Gasto Avulso", key=f"salvar_av_{fka}"):
        if desc and val > 0:
            metodo_salvar = "Saldo" if metodo_p_av == "Saldo em Conta" else "Cartao"
            periodo_fat_av = periodo_ativo
            if metodo_salvar == "Cartao":
                cartao_obj_av = obter_dados_cartao(dados, cartao_pagamento_av)
                if cartao_obj_av:
                    periodo_fat_av = calcular_periodo_fatura(str(data_gasto), cartao_obj_av.get("fechamento", 1))
            novo_avulso = {
                "data": str(data_gasto),
                "descricao": desc,
                "valor": val,
                "metodo_pagamento": metodo_salvar,
                "conta": conta_pagamento_av,
                "cartao_nome": cartao_pagamento_av,
                "periodo_fatura": periodo_fat_av
            }
            dados.setdefault("gastos_avulsos", []).append(novo_avulso)
            
            # Deduz ou adiciona na fatura
            if metodo_salvar == "Saldo":
                alterar_saldo(dados, conta_pagamento_av, val, "subtrair")
            else:
                alterar_fatura(dados, cartao_pagamento_av, val, "somar")
            
            salvar_dados_usuario(usuario_atual, dados)
            st.session_state.avulso_form_key += 1
            st.success("Gasto avulso adicionado!")
            st.rerun()
        else:
            st.error("Preencha a descrição e um valor maior que zero.")

    # Mostrar Gastos Avulsos
    gastos_avulsos_mes = [g for g in dados.get("gastos_avulsos", []) if pertence_periodo_cartao(g, periodo_ativo)]
    if gastos_avulsos_mes:
        st.markdown("#### Seus Gastos Avulsos Registrados")

        filtro_av = st.text_input("🔎 Buscar por descrição:", key="filtro_avulsos")

        for idx, item in enumerate(dados["gastos_avulsos"]):
            if not pertence_periodo_cartao(item, periodo_ativo):
                continue
            if filtro_av and filtro_av.lower() not in item["descricao"].lower():
                continue

            metodo_item = item.get("metodo_pagamento", "Saldo")
            local_item = item.get("conta", "Nu") if metodo_item == "Saldo" else item.get("cartao_nome", "Cartão Nu")

            col_d, col_v, col_del = st.columns([3, 2, 1])
            col_d.markdown(f"**{item['descricao']}**\n\n*{formatar_data_br(item['data'])} — {local_item}*")
            col_v.markdown(f"R$ {item['valor']:,.2f}")

            confirm_key = f"confirma_del_av_{idx}"
            if col_del.button("🗑️", key=f"del_av_{idx}", help="Excluir este lançamento"):
                st.session_state[confirm_key] = True

            if st.session_state.get(confirm_key):
                st.warning(f"Excluir '{item['descricao']}' (R$ {item['valor']:,.2f})? O valor volta para '{local_item}'.")
                c_sim, c_nao = st.columns(2)
                if c_sim.button("Sim, excluir", key=f"conf_sim_av_{idx}"):
                    if metodo_item == "Saldo":
                        alterar_saldo(dados, local_item, item["valor"], "somar")
                    else:
                        alterar_fatura(dados, local_item, item["valor"], "subtrair")
                    dados["gastos_avulsos"].pop(idx)
                    salvar_dados_usuario(usuario_atual, dados)
                    del st.session_state[confirm_key]
                    st.rerun()
                if c_nao.button("Cancelar", key=f"conf_nao_av_{idx}"):
                    del st.session_state[confirm_key]
                    st.rerun()
            st.markdown("---")

        with st.expander("⚠️ Apagar todo o histórico de avulsos deste mês"):
            st.caption("Atenção: isso remove todos os lançamentos do mês, sem estornar os saldos/faturas.")
            confirma_limpar_av = st.checkbox("Confirmo que quero apagar tudo", key="chk_limpar_av")
            if st.button("Limpar Histórico de Avulsos (Mês)", key="limpar_avulsos", disabled=not confirma_limpar_av):
                dados["gastos_avulsos"] = [g for g in dados.get("gastos_avulsos", []) if not pertence_periodo_cartao(g, periodo_ativo)]
                salvar_dados_usuario(usuario_atual, dados)
                st.rerun()

# ----------------- ABA 5: GASTOS DE TERCEIROS (CONTAS DE DEVEDORES POR PESSOA) -----------------
with abas[4]:
    st.markdown("<h3 class='titulo-secao'>👥 Amigos & Terceiros (Faturas & Empréstimos)</h3>", unsafe_allow_html=True)
    
    sub_abas_terceiros = st.tabs(["➕ Cadastrar Lançamento", "👤 Contas dos Devedores", "✅ Reembolsos Recebidos"])
    
    # SUB-ABA 5.1: CADASTRAR COMPRA OU EMPRÉSTIMO
    with sub_abas_terceiros[0]:
        opcao_terc = st.radio("Selecione o tipo de Lançamento de Terceiro:", ["Compras no Meu Cartão", "Empréstimos (PIX/Dinheiro)"], horizontal=True)
        
        if opcao_terc == "Compras no Meu Cartão":
            st.markdown("**Adicionar Compra de Terceiro no seu Cartão de Crédito**")
            st.caption("Aumenta a fatura do seu cartão. A pessoa fica te devendo mensalmente.")

            if "terc_form_key" not in st.session_state:
                st.session_state.terc_form_key = 0
            fkt = st.session_state.terc_form_key

            nome = st.text_input("Nome de quem comprou:", key=f"nome_terc_{fkt}")
            desc = st.text_input("Descrição da compra (Ex: Pizza, Tênis, Netflix dividido)", key=f"desc_terc_{fkt}")
            val = st.number_input("Valor da compra/parcela (R$)", min_value=0.0, step=10.0, format="%.2f", key=f"val_terc_{fkt}")

            # Selecionar cartão utilizado
            cartao_utilizado = st.selectbox("Qual Cartão foi Utilizado?", [c["nome"] for c in dados["cartoes"]], key=f"cartao_terc_{fkt}")

            tipo_lanc_terc = st.radio(
                "Tipo de Lançamento:",
                ["Única (só este mês)", "Parcelada (número fixo de parcelas)", "Recorrente (todo mês, até eu encerrar)"],
                key=f"tipo_terc_{fkt}", horizontal=True
            )
            st.caption("Use 'Recorrente' para gastos divididos todo mês, como uma assinatura de streaming.")

            total_parc_terc = 1
            if tipo_lanc_terc == "Parcelada (número fixo de parcelas)":
                total_parc_terc = st.number_input("Número de parcelas de terceiro:", min_value=2, max_value=48, value=2, step=1, key=f"parc_terc_{fkt}")

            pago_terc = st.checkbox("Marcar já como paga esta parcela/mês?", key=f"pago_terc_{fkt}")

            if st.button("Salvar Lançamento", key=f"salvar_terc_{fkt}"):
                if nome and desc and val > 0:
                    data_inicial = datetime(ano_selecionado, mes_num, 1)
                    nome_formatado = nome.title().strip()

                    if tipo_lanc_terc == "Recorrente (todo mês, até eu encerrar)":
                        novo_terc_rec = {
                            "id": f"tercrec_{datetime.now().strftime('%Y%m%d%H%M%S%f')}",
                            "nome": nome_formatado,
                            "descricao": desc,
                            "valor": val,
                            "cartao_nome": cartao_utilizado,
                            "inicio": periodo_ativo,
                            "fim": None,
                            "pagamentos": {periodo_ativo: pago_terc},
                            "valores_override": {}
                        }
                        dados.setdefault("terceiros_recorrentes", []).append(novo_terc_rec)
                        if pago_terc:
                            alterar_fatura(dados, cartao_utilizado, val, "somar")

                    elif tipo_lanc_terc == "Parcelada (número fixo de parcelas)":
                        lista_periodos = get_proximos_meses(data_inicial, total_parc_terc)
                        for i, periodo in enumerate(lista_periodos):
                            status_pago_parc = pago_terc if i == 0 else False
                            nova_compra = {
                                "data": f"{periodo}-01",
                                "nome": nome_formatado,
                                "descricao": f"{desc} (Parc. {i+1}/{total_parc_terc})",
                                "valor": val,
                                "pago": status_pago_parc,
                                "cartao_nome": cartao_utilizado,
                                "periodo_fatura": periodo
                            }
                            dados.setdefault("gastos_terceiros_cartao", []).append(nova_compra)

                            # Aumenta a fatura do cartão
                            if i == 0:
                                alterar_fatura(dados, cartao_utilizado, val, "somar")
                    else:
                        cartao_obj_terc = obter_dados_cartao(dados, cartao_utilizado)
                        periodo_fat_terc = calcular_periodo_fatura(f"{periodo_ativo}-01", cartao_obj_terc.get("fechamento", 1)) if cartao_obj_terc else periodo_ativo
                        nova_compra = {
                            "data": f"{periodo_ativo}-01",
                            "nome": nome_formatado,
                            "descricao": desc,
                            "valor": val,
                            "pago": pago_terc,
                            "cartao_nome": cartao_utilizado,
                            "periodo_fatura": periodo_fat_terc
                        }
                        dados.setdefault("gastos_terceiros_cartao", []).append(nova_compra)
                        alterar_fatura(dados, cartao_utilizado, val, "somar")

                    salvar_dados_usuario(usuario_atual, dados)
                    st.session_state.terc_form_key += 1
                    st.success(f"Lançamento de cartão registrado para {nome_formatado}! Fatura do '{cartao_utilizado}' atualizada.")
                    st.rerun()
                else:
                    st.error("Preencha nome, descrição e um valor maior que zero.")
                    
        else:
            st.markdown("**Adicionar Empréstimo direto em Dinheiro ou PIX**")
            st.caption("Retira dinheiro da sua conta corrente imediatamente para dar a um terceiro.")
            
            with st.form("form_emprestimo_terceiros", clear_on_submit=True):
                nome_emp = st.text_input("Quem pediu o empréstimo?")
                desc_emp = st.text_input("Motivo do empréstimo (Ex: Empréstimo para Luz)")
                val_emp = st.number_input("Valor Emprestado (R$)", min_value=0.0, step=10.0, format="%.2f")
                
                # Selecionar de qual conta tirar o dinheiro
                conta_origem_emp = st.selectbox("Retirar Dinheiro de qual Conta/Dinheiro?", [c["nome"] for c in dados["contas"]])
                
                enviar_e = st.form_submit_button("Registrar Empréstimo")
                if enviar_e and nome_emp and val_emp > 0:
                    nome_formatado_emp = nome_emp.title().strip()
                    novo_emp = {
                        "data": str(datetime.now().date()),
                        "nome": nome_formatado_emp,
                        "descricao": desc_emp,
                        "valor": val_emp,
                        "pago": False,
                        "conta_origem": conta_origem_emp,
                        "forma_recebimento": "",
                        "conta_destino": ""
                    }
                    dados.setdefault("gastos_terceiros_emprestimo", []).append(novo_emp)
                    
                    # Deduz da conta de origem automaticamente
                    alterar_saldo(dados, conta_origem_emp, val_emp, "subtrair")
                    
                    salvar_dados_usuario(usuario_atual, dados)
                    st.success(f"Empréstimo registrado para {nome_formatado_emp}! R$ {val_emp:,.2f} deduzidos da conta '{conta_origem_emp}'.")
                    st.rerun()

    # SUB-ABA 5.2: CONTAS DOS DEVEDORES (VISTA POR PESSOA)
    with sub_abas_terceiros[1]:
        st.markdown("**Gestão de Débitos Pendentes por Devedor**")
        st.caption("Consolidação automática de contas por devedor. Clique na pessoa para ver o detalhamento ou quitar débitos parciais ou totais.")
        
        # Pegar todos os registros pendentes de cartão e empréstimo
        unpaid_cartao = [t for t in dados.get("gastos_terceiros_cartao", []) if not t["pago"]]
        unpaid_emprestimo = [e for e in dados.get("gastos_terceiros_emprestimo", []) if not e["pago"]]
        recorrentes_terc_ativos = [
            r for r in dados.get("terceiros_recorrentes", [])
            if r["inicio"] <= periodo_ativo and (r["fim"] is None or periodo_ativo < r["fim"])
        ]
        
        # Encontrar todas as pessoas com débitos ativos
        nomes_devedores = sorted(list(set(
            [t["nome"] for t in unpaid_cartao] + 
            [e["nome"] for e in unpaid_emprestimo] +
            [r["nome"] for r in recorrentes_terc_ativos]
        )))
        
        if not nomes_devedores:
            st.success("🎉 Excelente! Ninguém te deve nada neste momento.")
        else:
            for nome_dev in nomes_devedores:
                # Filtrar itens específicos desse devedor
                debitos_cartao = [t for t in unpaid_cartao if t["nome"].lower() == nome_dev.lower()]
                debitos_emp = [e for e in unpaid_emprestimo if e["nome"].lower() == nome_dev.lower()]
                debitos_rec_terc = [r for r in recorrentes_terc_ativos if r["nome"].lower() == nome_dev.lower()]

                # Valor referente a ESTE mês (parcela do mês corrente) vs. total ainda pendente em todos os meses futuros
                debitos_cartao_mes = [t for t in debitos_cartao if pertence_periodo_cartao(t, periodo_ativo)]
                total_dev_cartao_mes = sum(t["valor"] for t in debitos_cartao_mes)
                total_dev_cartao_geral = sum(t["valor"] for t in debitos_cartao)
                total_dev_emp = sum(e["valor"] for e in debitos_emp)
                total_dev_rec_pendente = sum(
                    r["valores_override"].get(periodo_ativo, r["valor"])
                    for r in debitos_rec_terc if not r["pagamentos"].get(periodo_ativo, False)
                )
                total_dev_geral = total_dev_cartao_geral + total_dev_emp + total_dev_rec_pendente
                
                # Expandível customizado por devedor com balanço
                with st.expander(f"👤 {nome_dev} — Deve no total: R$ {total_dev_geral:,.2f}"):
                    st.markdown(f"**Balanço do Devedor:**")
                    st.markdown(f"- 💳 Valor do Cartão referente a {mes_selecionado}/{ano_selecionado}: **R$ {total_dev_cartao_mes:,.2f}**")
                    if total_dev_cartao_geral > total_dev_cartao_mes:
                        st.caption(f"Total ainda pendente no cartão, somando parcelas de outros meses: R$ {total_dev_cartao_geral:,.2f}")
                    st.markdown(f"- 💸 Débitos de Empréstimos (Pix/Dinheiro): **R$ {total_dev_emp:,.2f}**")
                    if debitos_rec_terc:
                        st.markdown(f"- 🔁 Recorrentes (referente a {mes_selecionado}/{ano_selecionado}): **R$ {total_dev_rec_pendente:,.2f}**")
                    st.markdown("---")
                    
                    # Mostrar Detalhamento do Cartão
                    if debitos_cartao:
                        st.markdown("**💳 Detalhamento de Compras no Cartão:**")
                        for idx_terc, t in enumerate(dados["gastos_terceiros_cartao"]):
                            if t not in debitos_cartao:
                                continue
                            desc_visual = t["descricao"]
                            base_desc = None
                            
                            # Parser de Parcelas Restantes
                            match = re.search(r"^(.*?)\s*\(Parc\.\s*(\d+)/(\d+)\)", desc_visual)
                            if match:
                                base_desc = match.group(1).strip()
                                total_parc = int(match.group(3))
                                
                                # Buscar todas as parcelas pendentes para essa pessoa
                                parcelas_totais = [g for g in dados.get("gastos_terceiros_cartao", []) if g["nome"].lower() == nome_dev.lower() and base_desc in g["descricao"]]
                                parcelas_pendentes = [g for g in parcelas_totais if not g["pago"]]
                                desc_visual = f"{t['descricao']} *(Faltam {len(parcelas_pendentes)} de {total_parc} parcelas)*"
                                
                            col_terc_txt, col_terc_btn1, col_terc_btn2 = st.columns([3, 1, 1])
                            col_terc_txt.write(f"- **{formatar_data_br(t['data'])}**: {desc_visual} — **R$ {t['valor']:,.2f}** *({t.get('cartao_nome', 'Cartão Nu')})*")

                            conv_key = f"conv_{idx_terc}"
                            if col_terc_btn1.button("🔁 Converter", key=f"btn_{conv_key}", help="Já paguei essa parcela com meu dinheiro; a pessoa passa a me dever essa parcela como empréstimo direto."):
                                st.session_state[conv_key] = True

                            del_key = f"delterc_{idx_terc}"
                            if col_terc_btn2.button("🗑️", key=f"btn_{del_key}", help="Excluir este lançamento (erro ou dívida perdoada)"):
                                st.session_state[del_key] = True

                            if st.session_state.get(conv_key):
                                # Converte SOMENTE esta parcela/lançamento (as demais parcelas, se houver, continuam no cartão)
                                itens_a_converter = [t]
                                valor_total_conv = sum(g["valor"] for g in itens_a_converter)

                                st.info(f"Você está convertendo '{t['descricao']}' — R$ {valor_total_conv:,.2f} — de dívida no cartão para empréstimo direto.")
                                conta_saida_conv = st.selectbox(
                                    "De qual conta sua saiu o dinheiro para pagar essa parcela agora?",
                                    [c["nome"] for c in dados["contas"] if c["tipo"] == "Normal"],
                                    key=f"conta_{conv_key}"
                                )
                                col_conv_sim, col_conv_nao = st.columns(2)
                                if col_conv_sim.button("Confirmar conversão", key=f"conf_{conv_key}"):
                                    alterar_fatura(dados, t.get("cartao_nome", ""), valor_total_conv, "subtrair")
                                    alterar_saldo(dados, conta_saida_conv, valor_total_conv, "subtrair")
                                    dados["gastos_terceiros_cartao"] = [g for g in dados["gastos_terceiros_cartao"] if g not in itens_a_converter]
                                    dados.setdefault("gastos_terceiros_emprestimo", []).append({
                                        "data": str(datetime.now().date()),
                                        "nome": nome_dev,
                                        "descricao": f"Convertido de compra no cartão: {t['descricao']}",
                                        "valor": valor_total_conv,
                                        "pago": False,
                                        "conta_origem": conta_saida_conv,
                                        "forma_recebimento": "",
                                        "conta_destino": ""
                                    })
                                    salvar_dados_usuario(usuario_atual, dados)
                                    del st.session_state[conv_key]
                                    st.success(f"Convertido! A fatura do cartão foi reduzida em R$ {valor_total_conv:,.2f} e {nome_dev} agora te deve isso como empréstimo direto.")
                                    st.rerun()
                                if col_conv_nao.button("Cancelar", key=f"canc_{conv_key}"):
                                    del st.session_state[conv_key]
                                    st.rerun()

                            if st.session_state.get(del_key):
                                st.warning(f"Excluir '{t['descricao']}' (R$ {t['valor']:,.2f})?")
                                motivo_del = st.radio(
                                    "Motivo:",
                                    ["Foi um erro de lançamento (estornar da fatura)", "Estou perdoando a dívida (a fatura já foi cobrada, manter como está)"],
                                    key=f"motivo_{del_key}"
                                )
                                col_del_sim, col_del_nao = st.columns(2)
                                if col_del_sim.button("Confirmar exclusão", key=f"conf_{del_key}"):
                                    if motivo_del.startswith("Foi um erro"):
                                        alterar_fatura(dados, t.get("cartao_nome", ""), t["valor"], "subtrair")
                                    dados["gastos_terceiros_cartao"].pop(idx_terc)
                                    salvar_dados_usuario(usuario_atual, dados)
                                    del st.session_state[del_key]
                                    st.success("Lançamento excluído.")
                                    st.rerun()
                                if col_del_nao.button("Cancelar", key=f"cancdel_{del_key}"):
                                    del st.session_state[del_key]
                                    st.rerun()
                            
                    # Mostrar Detalhamento do Empréstimo
                    if debitos_emp:
                        st.markdown("**💸 Detalhamento de Empréstimos (Dinheiro/Pix):**")
                        for idx_emp, e in enumerate(dados["gastos_terceiros_emprestimo"]):
                            if e not in debitos_emp:
                                continue
                            col_emp_txt, col_emp_btn = st.columns([4, 1])
                            col_emp_txt.write(f"- **{formatar_data_br(e['data'])}**: {e['descricao']} — **R$ {e['valor']:,.2f}** *(Retirado de: {e.get('conta_origem', 'Dinheiro')})*")

                            del_key_emp = f"delemp_{idx_emp}"
                            if col_emp_btn.button("🗑️", key=f"btn_{del_key_emp}", help="Excluir este empréstimo (erro ou dívida perdoada)"):
                                st.session_state[del_key_emp] = True

                            if st.session_state.get(del_key_emp):
                                st.warning(f"Excluir empréstimo '{e['descricao']}' (R$ {e['valor']:,.2f})?")
                                motivo_del_emp = st.radio(
                                    "Motivo:",
                                    ["Foi um erro de lançamento (devolver o valor para minha conta)", "Estou perdoando a dívida (o dinheiro já foi entregue, não devolver)"],
                                    key=f"motivo_{del_key_emp}"
                                )
                                col_del_sim_e, col_del_nao_e = st.columns(2)
                                if col_del_sim_e.button("Confirmar exclusão", key=f"conf_{del_key_emp}"):
                                    if motivo_del_emp.startswith("Foi um erro"):
                                        alterar_saldo(dados, e.get("conta_origem", ""), e["valor"], "somar")
                                    dados["gastos_terceiros_emprestimo"].pop(idx_emp)
                                    salvar_dados_usuario(usuario_atual, dados)
                                    del st.session_state[del_key_emp]
                                    st.success("Empréstimo excluído.")
                                    st.rerun()
                                if col_del_nao_e.button("Cancelar", key=f"cancdel_{del_key_emp}"):
                                    del st.session_state[del_key_emp]
                                    st.rerun()

                    # Mostrar Detalhamento dos Recorrentes de Terceiros
                    if debitos_rec_terc:
                        st.markdown("**🔁 Detalhamento de Recorrentes:**")
                        for rec_t in debitos_rec_terc:
                            valor_mes_rt = rec_t["valores_override"].get(periodo_ativo, rec_t["valor"])
                            pago_mes_rt = rec_t["pagamentos"].get(periodo_ativo, False)

                            col_rt1, col_rt2 = st.columns([3, 1])
                            col_rt1.write(f"- **{rec_t['descricao']}** 🔁 — **R$ {valor_mes_rt:,.2f}** *({rec_t.get('cartao_nome', 'Cartão Nu')})*")
                            novo_pago_rt = col_rt2.checkbox("Pago", value=pago_mes_rt, key=f"rt_pago_{rec_t['id']}_{periodo_ativo}")

                            if novo_pago_rt != pago_mes_rt:
                                if novo_pago_rt:
                                    alterar_fatura(dados, rec_t.get("cartao_nome", ""), valor_mes_rt, "somar")
                                else:
                                    alterar_fatura(dados, rec_t.get("cartao_nome", ""), valor_mes_rt, "subtrair")
                                rec_t["pagamentos"][periodo_ativo] = novo_pago_rt
                                salvar_dados_usuario(usuario_atual, dados)
                                st.rerun()

                            with st.expander(f"⚙️ Gerenciar '{rec_t['descricao']}' (recorrente)"):
                                novo_valor_rt = st.number_input(
                                    f"Valor de {mes_selecionado}/{ano_selecionado}:", min_value=0.0,
                                    value=float(valor_mes_rt), step=5.0, format="%.2f", key=f"editval_rt_{rec_t['id']}"
                                )
                                if st.button("Salvar novo valor deste mês", key=f"btn_editval_rt_{rec_t['id']}"):
                                    rec_t["valores_override"][periodo_ativo] = novo_valor_rt
                                    salvar_dados_usuario(usuario_atual, dados)
                                    st.rerun()
                                st.markdown("---")
                                if st.button(f"🛑 Encerrar a partir de {mes_selecionado}/{ano_selecionado}", key=f"encerrar_rt_{rec_t['id']}"):
                                    rec_t["fim"] = periodo_ativo
                                    salvar_dados_usuario(usuario_atual, dados)
                                    st.success("Recorrência encerrada a partir deste mês.")
                                    st.rerun()
                                st.markdown("---")
                                confirma_del_rt = st.checkbox("Confirmo que quero excluir esta recorrência por completo", key=f"chk_del_rt_{rec_t['id']}")
                                if st.button("🗑️ Excluir recorrência por completo", key=f"del_rt_{rec_t['id']}", disabled=not confirma_del_rt):
                                    dados["terceiros_recorrentes"] = [r for r in dados["terceiros_recorrentes"] if r["id"] != rec_t["id"]]
                                    salvar_dados_usuario(usuario_atual, dados)
                                    st.rerun()
                            
                    st.markdown("---")
                    st.markdown(f"**📥 Registrar Recebimento / Quitação de {nome_dev}**")
                    
                    # Criar opções para quitação parcial ou total de itens individuais
                    opcoes_debitos = []
                    for idx, t in enumerate(dados.get("gastos_terceiros_cartao", [])):
                        if not t["pago"] and t["nome"].lower() == nome_dev.lower():
                            opcoes_debitos.append({
                                "label": f"💳 Cartão: {t['descricao']} (R$ {t['valor']:,.2f})",
                                "tipo": "cartao",
                                "index": idx,
                                "valor": t["valor"],
                                "descricao": t["descricao"]
                            })
                    for idx, e in enumerate(dados.get("gastos_terceiros_emprestimo", [])):
                        if not e["pago"] and e["nome"].lower() == nome_dev.lower():
                            opcoes_debitos.append({
                                "label": f"💸 Empréstimo: {e['descricao']} (R$ {e['valor']:,.2f})",
                                "tipo": "emprestimo",
                                "index": idx,
                                "valor": e["valor"],
                                "descricao": e["descricao"]
                            })
                            
                    if opcoes_debitos:
                        item_labels = [opt["label"] for opt in opcoes_debitos]
                        item_sel_label = st.selectbox(f"Selecione qual item foi pago:", item_labels, key=f"sel_quitar_{nome_dev}")
                        
                        selected_opt = next(opt for opt in opcoes_debitos if opt["label"] == item_sel_label)
                        
                        with st.form(f"form_receber_pessoa_{nome_dev}_{selected_opt['tipo']}_{selected_opt['index']}", clear_on_submit=True):
                            st.write(f"**Valor Original Pendente:** R$ {selected_opt['valor']:,.2f}")
                            valor_pago = st.number_input("Quanto a pessoa pagou de fato? (R$):", min_value=0.0, value=selected_opt["valor"], step=10.0, format="%.2f", key=f"pago_val_{nome_dev}")
                            
                            # Conta que vai receber o dinheiro
                            conta_dest_reemb = st.selectbox("Depositar na Conta / Banco:", [c["nome"] for c in dados["contas"]], key=f"cont_dest_{nome_dev}")
                            forma_rec = st.selectbox("Forma de Pagamento:", ["PIX", "Dinheiro", "Transferência"], key=f"form_rec_{nome_dev}")
                            
                            confirmar_rec = st.form_submit_button("Confirmar Pagamento 📥")
                            if confirmar_rec:
                                tipo_item = selected_opt["tipo"]
                                idx_orig = selected_opt["index"]
                                valor_original = selected_opt["valor"]
                                
                                if valor_pago <= 0:
                                    st.error("O valor pago deve ser maior que zero!")
                                else:
                                    if tipo_item == "cartao":
                                        item_db = dados["gastos_terceiros_cartao"][idx_orig]
                                    else:
                                        item_db = dados["gastos_terceiros_emprestimo"][idx_orig]
                                        
                                    # CASO 1: PAGAMENTO INTEGRAL (EXATO)
                                    if abs(valor_pago - valor_original) < 0.01:
                                        item_db["pago"] = True
                                        item_db["conta_destino"] = conta_dest_reemb
                                        item_db["data_pagamento"] = str(datetime.now().date())
                                        if tipo_item == "emprestimo":
                                            item_db["forma_recebimento"] = forma_rec

                                        dados.setdefault("recebimentos_terceiros", []).append({
                                            "data": str(datetime.now().date()),
                                            "nome": nome_dev,
                                            "descricao": selected_opt["descricao"],
                                            "valor": valor_pago,
                                            "conta_destino": conta_dest_reemb
                                        })
                                        alterar_saldo(dados, conta_dest_reemb, valor_pago, "somar")
                                        salvar_dados_usuario(usuario_atual, dados)
                                        st.success(f"Quitação total realizada! R$ {valor_pago:,.2f} adicionados à conta '{conta_dest_reemb}'.")
                                        st.rerun()
                                        
                                    # CASO 2: PAGOU A MENOS (PAGAMENTO PARCIAL)
                                    elif valor_pago < valor_original:
                                        novo_valor_pendente = valor_original - valor_pago
                                        item_db["valor"] = novo_valor_pendente
                                        item_db["descricao"] = f"{item_db['descricao']} (Parcial R$ {valor_pago:,.2f} recebido)"

                                        dados.setdefault("recebimentos_terceiros", []).append({
                                            "data": str(datetime.now().date()),
                                            "nome": nome_dev,
                                            "descricao": selected_opt["descricao"],
                                            "valor": valor_pago,
                                            "conta_destino": conta_dest_reemb
                                        })
                                        alterar_saldo(dados, conta_dest_reemb, valor_pago, "somar")
                                        salvar_dados_usuario(usuario_atual, dados)
                                        st.success(f"Reembolso parcial! {nome_dev} pagou R$ {valor_pago:,.2f} e ainda restam R$ {novo_valor_pendente:,.2f} pendentes neste item.")
                                        st.rerun()
                                        
                                    # CASO 3: PAGOU A MAIS (SOBROU TROCO / EXCESSO)
                                    elif valor_pago > valor_original:
                                        item_db["pago"] = True
                                        item_db["conta_destino"] = conta_dest_reemb
                                        item_db["data_pagamento"] = str(datetime.now().date())
                                        if tipo_item == "emprestimo":
                                            item_db["forma_recebimento"] = forma_rec
                                            
                                        excesso = valor_pago - valor_original

                                        # Registra a parte "base" da dívida como recebimento de terceiro;
                                        # o excesso é lançado à parte como Receita (evita duplicar no extrato)
                                        dados.setdefault("recebimentos_terceiros", []).append({
                                            "data": str(datetime.now().date()),
                                            "nome": nome_dev,
                                            "descricao": selected_opt["descricao"],
                                            "valor": valor_original,
                                            "conta_destino": conta_dest_reemb
                                        })
                                        
                                        # Deposita valor total na conta
                                        alterar_saldo(dados, conta_dest_reemb, valor_pago, "somar")
                                        
                                        # Lança excesso como Receita
                                        nova_rec = {
                                            "data": str(datetime.now().date()),
                                            "descricao": f"Diferença paga a mais por {nome_dev} (Ref: {selected_opt['descricao']})",
                                            "valor": excesso,
                                            "conta": conta_dest_reemb
                                        }
                                        dados.setdefault("receitas", []).append(nova_rec)
                                        
                                        salvar_dados_usuario(usuario_atual, dados)
                                        st.success(f"Quitação total e Excesso registrado! Como o pagamento foi R$ {valor_pago:,.2f} (R$ {excesso:,.2f} a mais que a dívida), a diferença entrou como Receita e o saldo total foi para a conta '{conta_dest_reemb}'.")
                                        st.rerun()

    # SUB-ABA 5.3: HISTÓRICO DE REEMBOLSOS RECEBIDOS
    with sub_abas_terceiros[2]:
        st.markdown("**Histórico de Pagamentos e Reembolsos Concluídos**")
        
        pessoa_cartao_paga = [t for t in dados.get("gastos_terceiros_cartao", []) if t["pago"]]
        pessoa_emprestimo_pago = [e for e in dados.get("gastos_terceiros_emprestimo", []) if e["pago"]]
        
        if not pessoa_cartao_paga and not pessoa_emprestimo_pago:
            st.info("Nenhum reembolso completo registrado até o momento.")
        else:
            if pessoa_cartao_paga:
                st.markdown("**💳 Compras no Cartão (Pagas/Reembolsadas):**")
                for t in pessoa_cartao_paga:
                    st.write(f"- **{formatar_data_br(t['data'])}**: {t['nome']} pagou: **R$ {t['valor']:,.2f}** - *({t['descricao']})*")
            if pessoa_emprestimo_pago:
                st.markdown("**💸 Empréstimos (Quitados/Recebidos):**")
                for e in pessoa_emprestimo_pago:
                    st.write(f"- **{formatar_data_br(e['data'])}**: {e['nome']} quitou: **R$ {e['valor']:,.2f}** - *({e['descricao']})* -> Destino: {e.get('conta_destino', 'Nu')} ({e.get('forma_recebimento', 'PIX')})")

# ----------------- ABA 6: NOVA ABA SALDOS (CONTAS DINÂMICAS & FATURAS DE CARTÃO) -----------------
with abas[5]:
    st.markdown("<h3 class='titulo-secao'>💳 Gestão de Contas, Saldos & Cartões</h3>", unsafe_allow_html=True)
    st.caption("Crie, edite e acompanhe seus saldos bancários e faturas de cartões de crédito.")
    
    # 1. Visualizar Saldos do Mês Selecionado (baseado nas datas dos lançamentos, não no saldo "de hoje")
    st.markdown(f"#### Seus Saldos em {mes_selecionado}/{ano_selecionado}")
    st.caption("Cada mês mostra o saldo que a conta tinha NAQUELE mês — um ganho lançado em setembro não aparece em agosto.")
    
    col_c1, col_c2 = st.columns(2)
    with col_c1:
        st.markdown("**Contas Correntes (Dia a Dia):**")
        for conta in dados["contas"]:
            if conta.get("tipo", "Normal") == "Normal":
                saldo_mes_conta = calcular_saldo_conta_no_periodo(dados, conta["nome"], periodo_ativo)
                st.info(f"🏦 **{conta['nome']}:** R$ {saldo_mes_conta:,.2f}")
    with col_c2:
        st.markdown("**Poupança & Guardados:**")
        for conta in dados["contas"]:
            if conta.get("tipo", "Normal") == "Guardado":
                saldo_mes_conta = calcular_saldo_conta_no_periodo(dados, conta["nome"], periodo_ativo)
                st.success(f"🔒 **{conta['nome']}:** R$ {saldo_mes_conta:,.2f}")

    with st.expander("🔎 Ver saldo real de HOJE (para fazer transferências, pagamentos, etc.)"):
        st.caption("Este é o valor real disponível agora, independente do mês selecionado acima — use-o para saber quanto você realmente tem disponível hoje.")
        for conta in dados["contas"]:
            icone_hoje = "🔒" if conta.get("tipo") == "Guardado" else "🏦"
            st.write(f"{icone_hoje} **{conta['nome']}:** R$ {conta['saldo']:,.2f}")
                
    st.markdown("#### Faturas Ativas de Cartões de Crédito")
    col_cart_s = st.columns(2)
    for c_idx, cartao in enumerate(dados.get("cartoes", [])):
        with col_cart_s[c_idx % 2]:
            st.error(f"💳 **{cartao['nome']}:** R$ {cartao['fatura']:,.2f}\n\n*Fecha dia {cartao.get('fechamento', 1)} • Vence dia {cartao.get('vencimento', 10)}*")

    st.markdown("---")
    st.markdown(f"#### 📒 Extrato de {mes_selecionado}/{ano_selecionado}")
    st.caption("Saldo anterior (fim do mês passado) → movimentação do mês → saldo no fim do mês selecionado, conta por conta.")
    periodo_anterior_ativo = periodo_anterior(periodo_ativo)
    for conta in dados["contas"]:
        saldo_anterior_periodo = calcular_saldo_conta_no_periodo(dados, conta["nome"], periodo_anterior_ativo)
        saldo_final_periodo = calcular_saldo_conta_no_periodo(dados, conta["nome"], periodo_ativo)
        movimentacao = saldo_final_periodo - saldo_anterior_periodo
        icone_conta = "🔒" if conta.get("tipo") == "Guardado" else "🏦"
        col_ext1, col_ext2, col_ext3 = st.columns(3)
        col_ext1.metric(f"{icone_conta} {conta['nome']} — Saldo Anterior", f"R$ {saldo_anterior_periodo:,.2f}")
        col_ext2.metric("Movimentação no Mês", f"R$ {movimentacao:,.2f}")
        col_ext3.metric("Saldo no Fim do Mês", f"R$ {saldo_final_periodo:,.2f}")
                
    st.markdown("---")
    
    # Seções expansíveis para gerenciamento
    aba_ger_conta = st.tabs(["💸 Transferir", "💳 Pagar Fatura", "➕ Criar Conta/Cartão", "⚙️ Renomear/Excluir"])
    
    # 6.1 Lançamentos entre Contas (Transferências)
    with aba_ger_conta[0]:
        st.markdown("**Fazer Transferência entre Contas Bancárias**")
        with st.form("form_transferencia", clear_on_submit=True):
            contas_nomes = [c["nome"] for c in dados["contas"]]
            origem = st.selectbox("Conta de Origem (Sai o dinheiro):", contas_nomes)
            destino = st.selectbox("Conta de Destino (Entra o dinheiro):", contas_nomes)
            valor_transf = st.number_input("Valor da Transferência (R$)", min_value=0.0, step=50.0, format="%.2f")
            
            confirmar = st.form_submit_button("Realizar Transferência 🔁")
            if confirmar:
                if origem == destino:
                    st.error("A conta de origem não pode ser igual à conta de destino!")
                elif valor_transf <= 0:
                    st.error("O valor deve ser maior do que zero!")
                else:
                    if alterar_saldo(dados, origem, valor_transf, "subtrair"):
                        alterar_saldo(dados, destino, valor_transf, "somar")
                        
                        # Salva registro no histórico de transferências
                        nova_transf = {
                            "data": str(datetime.now().date()),
                            "origem": origem,
                            "destino": destino,
                            "valor": valor_transf
                        }
                        dados.setdefault("transferencias", []).append(nova_transf)
                        salvar_dados_usuario(usuario_atual, dados)
                        st.success(f"Transferência de R$ {valor_transf:,.2f} realizada com sucesso!")
                        st.rerun()
                    else:
                        st.error("Erro ao realizar transferência.")

    # 6.2 PAGAR FATURA DO CARTÃO DE CRÉDITO
    with aba_ger_conta[1]:
        st.markdown("**Pagar Fatura de Cartão de Crédito**")
        st.caption("Efetua o pagamento do saldo do cartão, reduzindo a fatura e retirando o dinheiro da sua conta bancária corrente.")
        
        with st.form("form_pagar_fatura", clear_on_submit=True):
            cartao_a_pagar = st.selectbox("Escolha qual Cartão quer pagar:", [c["nome"] for c in dados["cartoes"]])
            conta_pagadora = st.selectbox("Pagar usando o Saldo de qual Conta:", [c["nome"] for c in dados["contas"] if c["tipo"] == "Normal"])
            
            fatura_atual_val = next((c["fatura"] for c in dados["cartoes"] if c["nome"] == cartao_a_pagar), 0.0)
            st.write(f"Fatura total pendente deste cartão: **R$ {fatura_atual_val:,.2f}**")
            
            valor_pagamento_fat = st.number_input("Valor a pagar (R$):", min_value=0.0, value=fatura_atual_val, step=10.0, format="%.2f")
            
            pagar_fat_btn = st.form_submit_button("Confirmar Pagamento de Fatura 💳")
            if pagar_fat_btn and valor_pagamento_fat > 0:
                # Retirar do saldo da conta
                if alterar_saldo(dados, conta_pagadora, valor_pagamento_fat, "subtrair"):
                    # Reduzir fatura do cartão
                    alterar_fatura(dados, cartao_a_pagar, valor_pagamento_fat, "subtrair")
                    
                    # Salva nas transferências para registro histórico
                    dados.setdefault("transferencias", []).append({
                        "data": str(datetime.now().date()),
                        "origem": conta_pagadora,
                        "destino": f"Fatura {cartao_a_pagar}",
                        "valor": valor_pagamento_fat
                    })
                    salvar_dados_usuario(usuario_atual, dados)
                    st.success(f"Fatura do '{cartao_a_pagar}' no valor de R$ {valor_pagamento_fat:,.2f} foi paga usando a conta '{conta_pagadora}'!")
                    st.rerun()

    # 6.3 CRIAR CONTAS OU CARTÕES DINÂMICOS
    with aba_ger_conta[2]:
        col_c_c1, col_c_c2 = st.columns(2)
        
        with col_c_c1:
            st.markdown("**🏦 Criar Nova Conta Bancária**")
            with st.form("form_criar_conta_din", clear_on_submit=True):
                nome_nova_conta = st.text_input("Nome da Conta (Ex: Inter, Caixa, Carteira):")
                tipo_nova_conta = st.selectbox("Tipo da Conta:", ["Normal", "Guardado"])
                saldo_inicial = st.number_input("Saldo Inicial (R$):", min_value=0.0, step=10.0, format="%.2f")
                criar = st.form_submit_button("Criar Conta 🏦")
                if criar and nome_nova_conta:
                    ja_existe = any(c["nome"].lower() == nome_nova_conta.lower() for c in dados["contas"])
                    if ja_existe:
                        st.error("Já existe uma conta cadastrada com esse nome!")
                    else:
                        dados["contas"].append({
                            "nome": nome_nova_conta.strip(),
                            "saldo": saldo_inicial,
                            "tipo": tipo_nova_conta
                        })
                        salvar_dados_usuario(usuario_atual, dados)
                        st.success(f"Conta '{nome_nova_conta}' criada com sucesso!")
                        st.rerun()
                        
        with col_c_c2:
            st.markdown("**💳 Criar Novo Cartão de Crédito**")
            with st.form("form_criar_cartao_din", clear_on_submit=True):
                nome_novo_cartao = st.text_input("Nome do Cartão (Ex: Cartão Inter, Cartão Elo):")
                fatura_inicial = st.number_input("Fatura Inicial Pendente (R$):", min_value=0.0, step=10.0, format="%.2f")
                col_fech, col_venc = st.columns(2)
                dia_fechamento_novo = col_fech.number_input("Dia do Fechamento:", min_value=1, max_value=28, value=1, step=1)
                dia_vencimento_novo = col_venc.number_input("Dia do Vencimento:", min_value=1, max_value=28, value=10, step=1)
                st.caption("Compras feitas no dia do fechamento (ou depois) entram automaticamente na fatura do mês seguinte.")
                criar_c = st.form_submit_button("Criar Cartão 💳")
                if criar_c and nome_novo_cartao:
                    ja_existe = any(c["nome"].lower() == nome_novo_cartao.lower() for c in dados["cartoes"])
                    if ja_existe:
                        st.error("Já existe um cartão cadastrado com esse nome!")
                    else:
                        dados["cartoes"].append({
                            "nome": nome_novo_cartao.strip(),
                            "fatura": fatura_inicial,
                            "fechamento": int(dia_fechamento_novo),
                            "vencimento": int(dia_vencimento_novo)
                        })
                        salvar_dados_usuario(usuario_atual, dados)
                        st.success(f"Cartão '{nome_novo_cartao}' criado com sucesso!")
                        st.rerun()

    # 6.4 RENOMEAR OU EXCLUIR CONTAS / CARTÕES
    with aba_ger_conta[3]:
        st.markdown("**Renomear ou Excluir uma Conta ou Cartão**")
        tipo_ger = st.radio("Selecione qual deseja gerenciar:", ["Contas Bancárias", "Cartões de Crédito"], horizontal=True)
        
        if tipo_ger == "Contas Bancárias":
            conta_selecionada = st.selectbox("Selecione a Conta bancária para alterar:", [c["nome"] for c in dados["contas"]])
            col_ren, col_exc = st.columns(2)
            
            with col_ren:
                st.markdown("**Renomear Conta:**")
                novo_nome_conta = st.text_input("Novo Nome da Conta:", value=conta_selecionada)
                if st.button("Salvar Novo Nome"):
                    if novo_nome_conta and novo_nome_conta != conta_selecionada:
                        for c in dados["contas"]:
                            if c["nome"] == conta_selecionada:
                                c["nome"] = novo_nome_conta.strip()
                        # update other lists
                        for r in dados.get("receitas", []):
                            if r.get("conta") == conta_selecionada: r["conta"] = novo_nome_conta
                        for g in dados.get("gastos_fixos", []):
                            if g.get("conta") == conta_selecionada: g["conta"] = novo_nome_conta
                        for a in dados.get("gastos_avulsos", []):
                            if a.get("conta") == conta_selecionada: a["conta"] = novo_nome_conta
                        for e in dados.get("gastos_terceiros_emprestimo", []):
                            if e.get("conta_origem") == conta_selecionada: e["conta_origem"] = novo_nome_conta
                            if e.get("conta_destino") == conta_selecionada: e["conta_destino"] = novo_nome_conta
                        for t in dados.get("transferencias", []):
                            if t.get("origem") == conta_selecionada: t["origem"] = novo_nome_conta
                            if t.get("destino") == conta_selecionada: t["destino"] = novo_nome_conta
                        salvar_dados_usuario(usuario_atual, dados)
                        st.success("Conta renomeada!")
                        st.rerun()
            with col_exc:
                st.markdown("**Excluir Conta:**")
                st.warning("⚠️ Ao excluir a conta, seu saldo ativo será perdido.")
                confirma_exc_conta = st.checkbox(f"Confirmo que quero excluir '{conta_selecionada}'", key="chk_exc_conta")
                if st.button("Confirmar Exclusão ❌", disabled=not confirma_exc_conta):
                    if len(dados["contas"]) <= 1:
                        st.error("Você deve manter pelo menos uma conta ativa!")
                    else:
                        dados["contas"] = [c for c in dados["contas"] if c["nome"] != conta_selecionada]
                        salvar_dados_usuario(usuario_atual, dados)
                        st.success("Conta excluída!")
                        st.rerun()
                        
        else: # Cartões de Crédito
            cartao_selecionado = st.selectbox("Selecione o Cartão para gerenciar:", [c["nome"] for c in dados["cartoes"]])
            cartao_obj_sel = obter_dados_cartao(dados, cartao_selecionado)

            st.markdown("**Datas de Fechamento e Vencimento:**")
            col_fech_e, col_venc_e = st.columns(2)
            novo_fech = col_fech_e.number_input("Dia do Fechamento:", min_value=1, max_value=28, value=cartao_obj_sel.get("fechamento", 1), step=1, key="edit_fech")
            novo_venc = col_venc_e.number_input("Dia do Vencimento:", min_value=1, max_value=28, value=cartao_obj_sel.get("vencimento", 10), step=1, key="edit_venc")
            if st.button("Salvar Datas de Fatura"):
                cartao_obj_sel["fechamento"] = int(novo_fech)
                cartao_obj_sel["vencimento"] = int(novo_venc)
                salvar_dados_usuario(usuario_atual, dados)
                st.success("Datas de fechamento/vencimento atualizadas!")
                st.rerun()
            st.markdown("---")

            col_ren_c, col_exc_c = st.columns(2)
            
            with col_ren_c:
                st.markdown("**Renomear Cartão:**")
                novo_nome_cartao = st.text_input("Novo Nome do Cartão:", value=cartao_selecionado)
                if st.button("Salvar Novo Nome de Cartão"):
                    if novo_nome_cartao and novo_nome_cartao != cartao_selecionado:
                        for c in dados["cartoes"]:
                            if c["nome"] == cartao_selecionado:
                                c["nome"] = novo_nome_cartao.strip()
                        for g in dados.get("gastos_fixos", []):
                            if g.get("cartao_nome") == cartao_selecionado: g["cartao_nome"] = novo_nome_cartao
                        for a in dados.get("gastos_avulsos", []):
                            if a.get("cartao_nome") == cartao_selecionado: a["cartao_nome"] = novo_nome_cartao
                        for t in dados.get("gastos_terceiros_cartao", []):
                            if t.get("cartao_nome") == cartao_selecionado: t["cartao_nome"] = novo_nome_cartao
                        salvar_dados_usuario(usuario_atual, dados)
                        st.success("Cartão renomeado!")
                        st.rerun()
            with col_exc_c:
                st.markdown("**Excluir Cartão:**")
                st.warning("⚠️ Ao excluir o cartão, seu saldo de fatura ativa será deletado.")
                confirma_exc_cartao = st.checkbox(f"Confirmo que quero excluir '{cartao_selecionado}'", key="chk_exc_cartao")
                if st.button("Confirmar Exclusão de Cartão ❌", disabled=not confirma_exc_cartao):
                    if len(dados["cartoes"]) <= 1:
                        st.error("Você deve manter pelo menos um cartão ativo!")
                    else:
                        dados["cartoes"] = [c for c in dados["cartoes"] if c["nome"] != cartao_selecionado]
                        salvar_dados_usuario(usuario_atual, dados)
                        st.success("Cartão excluído!")
                        st.rerun()

# ----------------- ABA 7: POUPANÇA (GUARDADO) -----------------
with abas[6]:
    st.markdown("<h3 class='titulo-secao'>🔒 Dinheiro Guardado & Poupança</h3>", unsafe_allow_html=True)
    st.markdown("""
    <div class="card-didatico">
        💡 <b>Como funciona seu Dinheiro Guardado:</b><br>
        Toda conta que você cria na aba anterior e define como do tipo <b>"Guardado"</b> (como o Banco Pan ou as Caixinhas) 
        é listada aqui de forma integrada! Para guardar dinheiro, basta ir na aba anterior e fazer uma transferência de uma conta corrente normal para a sua conta de poupança.
    </div>
    """, unsafe_allow_html=True)
    
    # Calcular e mostrar saldos de investimentos (referentes ao mês/ano selecionado no topo)
    contas_guardado = [c for c in dados["contas"] if c.get("tipo", "Normal") == "Guardado"]
    saldos_guardado_mes = {c["nome"]: calcular_saldo_conta_no_periodo(dados, c["nome"], periodo_ativo) for c in contas_guardado}
    total_reservas = sum(saldos_guardado_mes.values())

    st.caption(f"Valores referentes a {mes_selecionado}/{ano_selecionado}.")
    st.markdown("#### Detalhamento das Economias:")
    for conta in contas_guardado:
        st.info(f"💰 **{conta['nome']}:** R$ {saldos_guardado_mes[conta['nome']]:,.2f}")
        
    st.markdown(f"### 📈 Total Acumulado Guardado: **R$ {total_reservas:,.2f}**")

    if contas_guardado:
        st.markdown("#### 📒 Extrato do Mês por Conta Guardada")
        periodo_anterior_guardado = periodo_anterior(periodo_ativo)
        for conta in contas_guardado:
            saldo_ant_g = calcular_saldo_conta_no_periodo(dados, conta["nome"], periodo_anterior_guardado)
            saldo_fim_g = calcular_saldo_conta_no_periodo(dados, conta["nome"], periodo_ativo)
            movimentacao_g = saldo_fim_g - saldo_ant_g
            col_g1, col_g2, col_g3 = st.columns(3)
            col_g1.metric(f"🔒 {conta['nome']} — Anterior", f"R$ {saldo_ant_g:,.2f}")
            col_g2.metric("Movimentação no Mês", f"R$ {movimentacao_g:,.2f}")
            col_g3.metric("Saldo no Fim do Mês", f"R$ {saldo_fim_g:,.2f}")

        with st.expander("🔎 Ver saldo real de HOJE (independente do mês selecionado)"):
            for conta in contas_guardado:
                st.write(f"🔒 **{conta['nome']}:** R$ {conta['saldo']:,.2f}")

    # Gráfico de Pizza: composição das reservas guardadas
    if contas_guardado and total_reservas > 0:
        st.markdown("#### 🥧 Composição das suas Reservas")
        df_pizza_guardado = pd.DataFrame([{"Conta": nome, "Valor": v} for nome, v in saldos_guardado_mes.items() if v > 0])
        fig_guardado = px.pie(df_pizza_guardado, values="Valor", names="Conta", hole=0.4,
                               color_discrete_sequence=px.colors.sequential.Purples_r)
        fig_guardado.update_layout(
            paper_bgcolor='rgba(0,0,0,0)',
            plot_bgcolor='rgba(0,0,0,0)',
            font_color="#FFFFFF",
            margin=dict(t=10, b=10, l=10, r=10),
            height=260
        )
        st.plotly_chart(fig_guardado, use_container_width=True)
    
    # Histórico de transferências para poupança (economias)
    transferencias_poupanca = []
    for t in dados.get("transferencias", []):
        dest_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["destino"]), "Normal")
        orig_tipo = next((c["tipo"] for c in dados["contas"] if c["nome"] == t["origem"]), "Normal")
        if dest_tipo == "Guardado" and orig_tipo == "Normal":
            transferencias_poupanca.append(t)
            
    if transferencias_poupanca:
        st.markdown("#### Histórico Recente de Economias Guardadas")
        df_dep = pd.DataFrame(transferencias_poupanca)
        df_dep = df_dep[["data", "origem", "destino", "valor"]]
        df_dep["data"] = df_dep["data"].apply(formatar_data_br)
        df_dep.columns = ["Data", "Origem", "Destino (Guardado)", "Valor Guardado (R$)"]
        st.dataframe(df_dep, use_container_width=True)

# ----------------- ABA 8: ALTERAÇÃO DE SENHA (PIN) -----------------
with abas[7]:
    st.markdown("<h3 class='titulo-secao'>⚙️ Alterar Sua Senha (PIN)</h3>", unsafe_allow_html=True)
    st.caption("Sua senha garante a privacidade dos seus dados financeiros se outras pessoas usarem o mesmo aplicativo.")
    
    with st.form("form_alterar_senha"):
        senha_atual = st.text_input("Digite o PIN atual (6 dígitos):", type="password", max_chars=6)
        nova_senha = st.text_input("Digite o NOVO PIN (6 dígitos numéricos):", type="password", max_chars=6)
        nova_senha_conf = st.text_input("Confirme o NOVO PIN (6 dígitos numéricos):", type="password", max_chars=6)
        
        atualizar_btn = st.form_submit_button("Alterar Senha")
        if atualizar_btn:
            if senha_atual != dados.get("pin", "123456"):
                st.error("PIN atual incorreto!")
            elif not nova_senha.isdigit() or len(nova_senha) != 6:
                st.error("O novo PIN deve conter exatamente 6 números.")
            elif nova_senha != nova_senha_conf:
                st.error("A confirmação da nova senha está diferente!")
            else:
                dados["pin"] = nova_senha
                salvar_dados_usuario(usuario_atual, dados)
                st.success("Sua senha (PIN) de 6 dígitos foi alterada com sucesso!")
                st.rerun()

    st.markdown("---")
    st.markdown("### 🔑 Pergunta de Segurança (Recuperação de PIN)")
    if dados.get("pergunta_seguranca"):
        st.caption(f"Pergunta atual: **{dados['pergunta_seguranca']}**")
    else:
        st.caption("Você ainda não configurou uma pergunta de segurança. Configure uma para poder recuperar seu PIN caso esqueça.")

    pergunta_escolhida_cfg = st.selectbox("Pergunta de Segurança:", PERGUNTAS_SEGURANCA_PADRAO, key="pergunta_cfg")
    pergunta_final_cfg = pergunta_escolhida_cfg
    if pergunta_escolhida_cfg == PERGUNTAS_SEGURANCA_PADRAO[-1]:
        pergunta_final_cfg = st.text_input("Digite sua pergunta personalizada:", key="pergunta_custom_cfg")
    resposta_cfg = st.text_input("Resposta:", key="resposta_cfg")
    if st.button("Salvar Pergunta de Segurança"):
        if pergunta_final_cfg and resposta_cfg:
            dados["pergunta_seguranca"] = pergunta_final_cfg
            dados["resposta_hash"] = hash_resposta(resposta_cfg)
            salvar_dados_usuario(usuario_atual, dados)
            st.success("Pergunta de segurança salva com sucesso!")
            st.rerun()
        else:
            st.error("Preencha a pergunta e a resposta.")
